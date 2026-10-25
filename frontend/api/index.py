import io
import logging
import os
import re
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import List, Optional
from zoneinfo import ZoneInfo

import bcrypt
import jwt
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from dotenv import load_dotenv
from fastapi import APIRouter, Depends, FastAPI, HTTPException, Query
from fastapi.responses import StreamingResponse
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field
from starlette.middleware.cors import CORSMiddleware

import seed_data
import recurrence_seed
import unicodedata

ROOT_DIR = Path(__file__).parent
# Caricamento variabili d'ambiente robusto per i due contesti:
# - Anteprima Emergent (supervisor): il file è in /app/backend/.env
# - Vercel (serverless): le variabili arrivano dalla dashboard (os.environ già popolato)
for _cand in [ROOT_DIR / ".env", ROOT_DIR.parent.parent / "backend" / ".env"]:
    if _cand.exists():
        load_dotenv(_cand)

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(name)s - %(levelname)s - %(message)s")
logger = logging.getLogger("agendavisite")

mongo_url = os.environ["MONGO_URL"]
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ["DB_NAME"]]

JWT_SECRET = os.environ["JWT_SECRET"]
JWT_ALGO = "HS256"
ACCESS_TOKEN_DAYS = int(os.environ.get("ACCESS_TOKEN_DAYS", "30"))
ROME = ZoneInfo("Europe/Rome")
VISIT_THRESHOLD_DAYS = 21
DUE_SOON_DAYS = 15  # finestra per considerare un incasso differito "in scadenza"

# Payment modes selectable per single order. days=None means no future suspension.
PAYMENT_MODES = {
    # Bonifici con scadenza -> promemoria sospeso
    "bonifico_30": {"label": "Bonifico bancario 30 giorni", "days": 30},
    "bonifico_60": {"label": "Bonifico bancario 60 giorni", "days": 60},
    "bonifico_60_fm": {"label": "Bonifico bancario 60 gg fine mese", "days": 60},
    "bonifico_90": {"label": "Bonifico bancario 90 giorni", "days": 90},
    "bonifico_vista_fattura": {"label": "Bonifico bancario vista fattura", "days": 15},
    "bonifico_anticipato": {"label": "Bonifico bancario anticipato", "days": None},
    # Pagamenti alla consegna / immediati -> nessun promemoria
    "pos_consegna": {"label": "Pagamento alla consegna tramite POS", "days": None},
    "contanti_consegna": {"label": "Contanti alla consegna", "days": None},
    "anticipato": {"label": "Anticipato", "days": None},
    "contrassegno": {"label": "Contrassegno", "days": None},
    # RID, Ri.Ba., Rifatturazione -> nessun promemoria
    "rid_60": {"label": "R.I.D. 60 giorni", "days": None},
    "riba_30": {"label": "Ri.Ba. 30 giorni", "days": None},
    "riba_60": {"label": "Ri.Ba. 60 giorni", "days": None},
    "riba_90": {"label": "Ri.Ba. 90 giorni", "days": None},
    "rifatturazione_pac": {"label": "Rifatturazione Pac", "days": None},
    # Rimesse dirette mezzo agente -> promemoria "da incassare"
    "agente_30": {"label": "Rimessa diretta 30 giorni mezzo agente", "days": 30},
    "agente_60": {"label": "Rimessa diretta 60 giorni mezzo agente", "days": 60},
    # Titolo allo scarico -> nessun promemoria
    "titolo_vista": {"label": "Titolo allo scarico a vista", "days": None},
    "titolo_30": {"label": "Titolo allo scarico 30 giorni", "days": None},
    "titolo_60": {"label": "Titolo allo scarico 60 giorni", "days": None},
    "titolo_90": {"label": "Titolo allo scarico 90 giorni", "days": None},
    # storico (non più selezionabile)
    "agente_90": {"label": "Pagamento mezzo Agente 90 giorni", "days": 90},
}
PAYMENT_MODE_ORDER = [
    "bonifico_30", "bonifico_60", "bonifico_60_fm", "bonifico_90", "bonifico_vista_fattura", "bonifico_anticipato",
    "agente_30", "agente_60",
    "pos_consegna", "contanti_consegna", "anticipato", "contrassegno",
    "rid_60", "riba_30", "riba_60", "riba_90", "rifatturazione_pac",
    "titolo_vista", "titolo_30", "titolo_60", "titolo_90",
]


def payment_mode_from_text(text) -> Optional[str]:
    """Riconduce la descrizione di pagamento di GooseApp alla modalità interna."""
    t = re.sub(r"[^a-z0-9 ]", " ", (str(text or "")).lower().replace("ri.ba", "riba").replace("r.i.d", "rid")
               .replace("b.b.", "bonifico ").replace("d.f.f.m", "fine mese").replace("f.m.", "fine mese")
               .replace("r.d.", "rimessa ").replace("r.b.", "riba "))
    t = re.sub(r"\brb\b", "riba", re.sub(r"\brd\b", "rimessa", t))
    t = " " + re.sub(r"\s+", " ", t).strip() + " "
    if not t.strip():
        return None
    if t.strip() in PAYMENT_MODES:
        return t.strip()
    days = next((d for d in ("90", "60", "30") if f" {d} " in t or f" {d}gg" in t or f" {d}g " in t), None)
    if "titolo" in t:
        return {"30": "titolo_30", "60": "titolo_60", "90": "titolo_90"}.get(days, "titolo_vista")
    if " riba " in t or "ricevuta bancaria" in t:
        return {"60": "riba_60", "90": "riba_90"}.get(days, "riba_30")
    if " rid " in t:
        return "rid_60"
    if "rifattur" in t:
        return "rifatturazione_pac"
    if "rimessa" in t or "agente" in t:
        return {"60": "agente_60", "90": "agente_90"}.get(days, "agente_30")
    if "bonifico" in t or " bb " in t:
        if "anticipat" in t:
            return "bonifico_anticipato"
        if "vista" in t:
            return "bonifico_vista_fattura"
        if days == "60" and ("fine mese" in t or " fm " in t or " dffm " in t):
            return "bonifico_60_fm"
        return {"30": "bonifico_30", "60": "bonifico_60", "90": "bonifico_90"}.get(days)
    if " pos " in t:
        return "pos_consegna"
    if "contant" in t:
        return "contanti_consegna"
    if "contrassegno" in t:
        return "contrassegno"
    if "anticipat" in t:
        return "anticipato"
    return None

app = FastAPI(title="Rodomisto Rappresentanze API")
api = APIRouter(prefix="/api")


@api.get("/health")
async def health():
    """Diagnostica: conferma che la funzione API risponde (nessuna auth)."""
    try:
        await db.command("ping")
        db_ok = True
    except Exception:
        db_ok = False
    return {"ok": True, "db": db_ok, "service": "rodomisto-api"}
bearer = HTTPBearer(auto_error=False)


@app.get("/health")
async def health():
    return {"status": "ok"}


@api.get("/health")
async def api_health():
    return {"status": "ok"}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def now_utc() -> datetime:
    return datetime.now(timezone.utc)


# --- Provincia: normalizzazione a sigla di 2 lettere maiuscole ------------
# Mappa nome esteso (lowercase, senza accenti) -> sigla. Copre tutte le
# province italiane. Se l'input è già una sigla di 2 lettere la lascia (upper).
PROVINCE_MAP = {
    "agrigento": "AG", "alessandria": "AL", "ancona": "AN", "aosta": "AO",
    "arezzo": "AR", "ascoli piceno": "AP", "asti": "AT", "avellino": "AV",
    "bari": "BA", "barletta-andria-trani": "BT", "barletta andria trani": "BT",
    "belluno": "BL", "benevento": "BN", "bergamo": "BG", "biella": "BI",
    "bologna": "BO", "bolzano": "BZ", "brescia": "BS", "brindisi": "BR",
    "cagliari": "CA", "caltanissetta": "CL", "campobasso": "CB",
    "carbonia-iglesias": "CI", "caserta": "CE", "catania": "CT",
    "catanzaro": "CZ", "chieti": "CH", "como": "CO", "cosenza": "CS",
    "cremona": "CR", "crotone": "KR", "cuneo": "CN", "enna": "EN",
    "fermo": "FM", "ferrara": "FE", "firenze": "FI", "foggia": "FG",
    "forli-cesena": "FC", "forli cesena": "FC", "frosinone": "FR",
    "genova": "GE", "gorizia": "GO", "grosseto": "GR", "imperia": "IM",
    "isernia": "IS", "la spezia": "SP", "l'aquila": "AQ", "laquila": "AQ",
    "latina": "LT", "lecce": "LE", "lecco": "LC", "livorno": "LI",
    "lodi": "LO", "lucca": "LU", "macerata": "MC", "mantova": "MN",
    "massa-carrara": "MS", "massa carrara": "MS", "matera": "MT",
    "messina": "ME", "milano": "MI", "modena": "MO", "monza e brianza": "MB",
    "monza e della brianza": "MB", "napoli": "NA", "novara": "NO",
    "nuoro": "NU", "oristano": "OR", "padova": "PD", "palermo": "PA",
    "parma": "PR", "pavia": "PV", "perugia": "PG", "pesaro e urbino": "PU",
    "pescara": "PE", "piacenza": "PC", "pisa": "PI", "pistoia": "PT",
    "pordenone": "PN", "potenza": "PZ", "prato": "PO", "ragusa": "RG",
    "ravenna": "RA", "reggio calabria": "RC", "reggio di calabria": "RC",
    "reggio emilia": "RE", "reggio nell'emilia": "RE", "rieti": "RI",
    "rimini": "RN", "roma": "RM", "rovigo": "RO", "salerno": "SA",
    "sassari": "SS", "savona": "SV", "siena": "SI", "siracusa": "SR",
    "sondrio": "SO", "taranto": "TA", "teramo": "TE", "terni": "TR",
    "torino": "TO", "trapani": "TP", "trento": "TN", "treviso": "TV",
    "trieste": "TS", "udine": "UD", "varese": "VA", "venezia": "VE",
    "verbano-cusio-ossola": "VB", "verbania": "VB", "vercelli": "VC",
    "verona": "VR", "vibo valentia": "VV", "vicenza": "VI", "viterbo": "VT",
}
_VALID_SIGLE = set(PROVINCE_MAP.values())


def normalize_provincia(raw: Optional[str]) -> str:
    """Restituisce sempre la sigla di 2 lettere MAIUSCOLE della provincia.
    - '' se vuoto
    - sigla nota (es. 'cz') -> 'CZ'
    - nome esteso (es. 'Catanzaro') -> 'CZ'
    - fallback: primi 2 caratteri alfabetici in maiuscolo."""
    s = (raw or "").strip()
    if not s:
        return ""
    low = s.lower().strip(". ")
    # accenti via rimozione semplice
    low = (low.replace("à", "a").replace("è", "e").replace("é", "e")
              .replace("ì", "i").replace("ò", "o").replace("ù", "u"))
    if len(s) == 2 and s.isalpha() and s.upper() in _VALID_SIGLE:
        return s.upper()
    if low in PROVINCE_MAP:
        return PROVINCE_MAP[low]
    if len(s) == 2 and s.isalpha():
        return s.upper()
    letters = "".join(ch for ch in s if ch.isalpha())
    return letters[:2].upper()


def parse_activity_date(raw: Optional[str], default: datetime) -> datetime:
    """Converte una data ISO (o 'YYYY-MM-DD') in datetime UTC. Usa il giorno
    locale Rome a mezzogiorno per evitare slittamenti di fuso. Se non valida,
    ritorna il default."""
    if not raw:
        return default
    try:
        txt = raw.strip()
        if len(txt) == 10 and txt[4] == "-" and txt[7] == "-":
            y, m, d = int(txt[:4]), int(txt[5:7]), int(txt[8:10])
            local = datetime(y, m, d, 12, 0, 0, tzinfo=ROME)
            return local.astimezone(timezone.utc)
        dt = datetime.fromisoformat(txt.replace("Z", "+00:00"))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt
    except (ValueError, TypeError):
        return default


def start_of_today_utc() -> datetime:
    local = datetime.now(ROME)
    start_local = local.replace(hour=0, minute=0, second=0, microsecond=0)
    return start_local.astimezone(timezone.utc)


def start_of_month_utc() -> datetime:
    """Mezzanotte (ora di Roma) del 1° giorno del mese corrente, in UTC.
    I contatori DA VISITARE / GESTITI si azzerano a questo istante ogni mese."""
    local = datetime.now(ROME)
    start_local = local.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    return start_local.astimezone(timezone.utc)


def hash_pw(pw: str) -> str:
    return bcrypt.hashpw(pw.encode("utf-8")[:72], bcrypt.gensalt()).decode("utf-8")


def verify_pw(pw: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(pw.encode("utf-8")[:72], hashed.encode("utf-8"))
    except Exception:
        return False


def create_token(username: str) -> str:
    now = now_utc()
    payload = {"sub": username, "iat": now, "exp": now + timedelta(days=ACCESS_TOKEN_DAYS)}
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGO)


def iso(dt: Optional[datetime]) -> Optional[str]:
    if dt is None:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc).isoformat()


def parse_day_bound(s: Optional[str], end: bool = False) -> Optional[datetime]:
    """Converte 'YYYY-MM-DD' (giorno locale Europe/Rome) nel relativo istante UTC.
    end=False -> inizio giornata (00:00); end=True -> inizio del giorno successivo
    (limite superiore esclusivo). Ritorna None se la stringa è vuota/non valida."""
    if not s:
        return None
    try:
        d = datetime.strptime(s.strip(), "%Y-%m-%d")
    except ValueError:
        return None
    local = d.replace(hour=0, minute=0, second=0, microsecond=0, tzinfo=ROME)
    if end:
        local = local + timedelta(days=1)
    return local.astimezone(timezone.utc)


async def get_current_user(creds: Optional[HTTPAuthorizationCredentials] = Depends(bearer)):
    unauth = HTTPException(status_code=401, detail="Credenziali non valide o scadute",
                           headers={"WWW-Authenticate": "Bearer"})
    if creds is None or creds.scheme.lower() != "bearer":
        raise unauth
    try:
        payload = jwt.decode(creds.credentials, JWT_SECRET, algorithms=[JWT_ALGO],
                             options={"require": ["sub", "exp"]})
        username = payload["sub"]
    except Exception:
        raise unauth
    user = await db.users.find_one({"username": username, "is_active": True}, {"_id": 0})
    if not user:
        raise unauth
    return user


# ---------------------------------------------------------------------------
# Models
# ---------------------------------------------------------------------------
class LoginRequest(BaseModel):
    username: str
    password: str


class Company(BaseModel):
    name: str


class CompanyUpdate(BaseModel):
    name: Optional[str] = None
    active: Optional[bool] = None


class GiroCreate(BaseModel):
    name: str
    localities: List[str] = Field(default_factory=list)


class GiroUpdate(BaseModel):
    name: Optional[str] = None
    localities: Optional[List[str]] = None
    active: Optional[bool] = None


class Site(BaseModel):
    """Sede di scarico aggiuntiva del cliente. Ogni sede può avere un proprio
    Giro territoriale e una propria Posizione nel giro (ordine di marcia)."""
    id: Optional[str] = None
    nominativo: str = ""
    indirizzo: str = ""
    cap: str = ""
    citta: str = ""
    provincia: str = ""
    zona: str = ""
    giro_id: Optional[str] = None
    position: Optional[int] = None
    goose_code: str = ""  # codice/riferimento sede su GooseApp (per il match ordini)


class ClientCreate(BaseModel):
    ragione_sociale: str
    provincia: str = ""
    giro_id: Optional[str] = None
    position: Optional[int] = None
    citta: str = ""
    zona: str = ""
    indirizzo: str = ""
    cap: str = ""
    telefono: str = ""
    email: str = ""
    partita_iva: str = ""
    sede_nome: str = ""  # nominativo della sede principale (opzionale)
    goose_code: str = ""  # codice sede principale su GooseApp
    sites: Optional[List[Site]] = None  # sedi di scarico aggiuntive
    agent: Optional[str] = None


class ClientUpdate(BaseModel):
    ragione_sociale: Optional[str] = None
    provincia: Optional[str] = None
    giro_id: Optional[str] = None
    position: Optional[int] = None
    citta: Optional[str] = None
    zona: Optional[str] = None
    indirizzo: Optional[str] = None
    cap: Optional[str] = None
    telefono: Optional[str] = None
    email: Optional[str] = None
    partita_iva: Optional[str] = None
    sede_nome: Optional[str] = None
    goose_code: Optional[str] = None
    sites: Optional[List[Site]] = None
    permanent_note: Optional[str] = None
    agent: Optional[str] = None


class EventCreate(BaseModel):
    client_id: str
    type: str  # order | collection | reschedule | note | suspension | visit(legacy)
    company_id: Optional[str] = None
    note_text: Optional[str] = None
    reschedule_days: Optional[int] = None
    reschedule_date: Optional[str] = None  # ISO date string
    payment_mode: Optional[str] = None  # for orders
    activity_date: Optional[str] = None  # data attività (ISO o YYYY-MM-DD); default oggi
    collection_method: Optional[str] = None  # per incassi: contanti | bonifico | assegno
    collection_ref_date: Optional[str] = None  # per bonifico: data del bonifico (ISO)
    site_id: Optional[str] = None  # sede di scarico selezionata ("principal" o id sede)
    site_nominativo: Optional[str] = None
    site_indirizzo: Optional[str] = None


# ---------------------------------------------------------------------------
# Serializers
# ---------------------------------------------------------------------------
def normalize_sites(sites) -> list:
    """Normalizza una lista di sedi aggiuntive (oggetti Site o dict)."""
    out = []
    for s in sites or []:
        d = s.dict() if hasattr(s, "dict") else dict(s)
        out.append({
            "id": d.get("id") or str(uuid.uuid4()),
            "nominativo": (d.get("nominativo") or "").strip(),
            "indirizzo": (d.get("indirizzo") or "").strip(),
            "cap": (d.get("cap") or "").strip(),
            "citta": (d.get("citta") or "").strip(),
            "provincia": normalize_provincia(d.get("provincia") or ""),
            "zona": (d.get("zona") or "").strip(),
            "giro_id": d.get("giro_id") or None,
            "position": d.get("position"),
            "goose_code": (d.get("goose_code") or "").strip(),
        })
    return out


def client_public(doc: dict) -> dict:
    extra = doc.get("extra") or {}
    return {
        "id": doc["id"],
        "ragione_sociale": doc.get("ragione_sociale", ""),
        "provincia": doc.get("provincia", ""),
        "giro_id": doc.get("giro_id"),
        "position": doc.get("position"),
        "citta": doc.get("citta", ""),
        "zona": doc.get("zona", ""),
        "indirizzo": doc.get("indirizzo", ""),
        "cap": doc.get("cap", ""),
        "telefono": doc.get("telefono", ""),
        "email": doc.get("email", ""),
        "agent": doc.get("agent", ""),
        "permanent_note": doc.get("permanent_note", ""),
        "last_visit_at": iso(doc.get("last_visit_at")),
        "snoozed_until": iso(doc.get("snoozed_until")),
        "partita_iva": extra.get("partita_iva", ""),
        "sede_nome": extra.get("sede_nome", ""),
        "goose_code": extra.get("goose_code", ""),
        "sites": extra.get("sites", []),
        "extra": extra,
    }


def event_public(doc: dict) -> dict:
    return {
        "id": doc["id"],
        "client_id": doc["client_id"],
        "type": doc["type"],
        "company_id": doc.get("company_id"),
        "company_name": doc.get("company_name"),
        "note_text": doc.get("note_text"),
        "reschedule_until": iso(doc.get("reschedule_until")),
        "agent": doc.get("agent"),
        "created_at": iso(doc.get("created_at")),
        "recurrence_company": doc.get("recurrence_company"),
        "recurrence_period": doc.get("recurrence_period"),
        "payment_mode": doc.get("payment_mode"),
        "payment_mode_label": PAYMENT_MODES.get(doc.get("payment_mode"), {}).get("label") if doc.get("payment_mode") else None,
        "due_at": iso(doc.get("due_at")),
        "collection_method": doc.get("collection_method"),
        "collection_ref_date": doc.get("collection_ref_date"),
        "site_id": doc.get("site_id"),
        "site_nominativo": doc.get("site_nominativo"),
        "site_indirizzo": doc.get("site_indirizzo"),
        "order_number": doc.get("order_number"),
        "ship_label": doc.get("ship_label"),
        "source": doc.get("source"),
    }


# ---------------------------------------------------------------------------
# Seeding / import
# ---------------------------------------------------------------------------
async def seed():
    for i, name in enumerate(seed_data.COMPANIES):
        existing = await db.companies.find_one({"name": name})
        if not existing:
            await db.companies.insert_one({
                "id": str(uuid.uuid4()), "name": name, "active": True,
                "order": i, "created_at": now_utc(),
            })

    giro_count = await db.giri.count_documents({})
    if giro_count == 0:
        for i, g in enumerate(seed_data.GIRI):
            await db.giri.insert_one({
                "id": str(uuid.uuid4()), "name": g["name"], "localities": g["localities"],
                "order": i, "active": True, "created_at": now_utc(),
            })

    client_count = await db.clients.count_documents({})
    if client_count == 0:
        giri_docs = await db.giri.find({}, {"_id": 0}).to_list(1000)
        loc_index = seed_data.build_locality_index(giri_docs)
        parsed = seed_data.parse_clients()
        docs = []
        deleted = 0
        for c in parsed:
            if seed_data.should_delete(c["ragione_sociale"], c["citta"]):
                deleted += 1
                continue
            giro_id, position = seed_data.resolve_assignment(
                c["ragione_sociale"], c["zona"], giri_docs, loc_index
            )
            docs.append({
                "id": str(uuid.uuid4()),
                "ragione_sociale": c["ragione_sociale"],
                "codice_azienda": c["codice_azienda"],
                "provincia": c["provincia"],
                "giro_id": giro_id,
                "position": position if position is not None else 999,
                "citta": c["citta"],
                "zona": c["zona"],
                "indirizzo": c["indirizzo"],
                "cap": c["cap"],
                "telefono": c["telefono"],
                "email": c["email"],
                "agent": c["agent"],
                "permanent_note": "",
                "last_visit_at": None,
                "snoozed_until": None,
                "extra": c["extra"],
                "created_by": c["agent"],
                "deleted_at": None,
                "created_at": now_utc(),
            })
        # Extra clients not present in the Excel.
        for ec in seed_data.EXTRA_CLIENTS:
            giro_id, position = seed_data.giro_position(giri_docs, ec["giro_name"], ec["locality"])
            docs.append({
                "id": str(uuid.uuid4()),
                "ragione_sociale": ec["ragione_sociale"],
                "codice_azienda": "",
                "provincia": ec.get("provincia", ""),
                "giro_id": giro_id,
                "position": position if position is not None else 999,
                "citta": ec.get("citta", ""),
                "zona": ec.get("zona", ""),
                "indirizzo": ec.get("indirizzo", ""),
                "cap": ec.get("cap", ""),
                "telefono": ec.get("telefono", ""),
                "email": ec.get("email", ""),
                "agent": ec.get("agent", "umberto"),
                "permanent_note": "",
                "last_visit_at": None,
                "snoozed_until": None,
                "extra": {},
                "created_by": ec.get("agent", "umberto"),
                "deleted_at": None,
                "created_at": now_utc(),
            })
        if docs:
            await db.clients.insert_many(docs)
        assigned = sum(1 for d in docs if d["giro_id"])
        logger.info("Imported %d clients (%d assigned, %d da verificare, %d deleted, %d extra)",
                    len(docs), assigned, len(docs) - assigned, deleted, len(seed_data.EXTRA_CLIENTS))

    await seed_recurrences()
    await import_clients_ott26()


async def import_clients_ott26():
    """Import idempotente (una tantum) dell'anagrafica 'clienti menu e pellegrini
    ott 26'. Match su P.IVA (poi codice azienda / ragione sociale): i clienti
    esistenti vengono aggiornati (agente, P.IVA, CF, contatti mancanti), i nuovi
    inseriti. Il giro è ricavato da 'Descrizione Zona' tra i giri dell'agente;
    se non univoco il cliente finisce in 'Da verificare'."""
    key = "import_clienti_menu_pellegrini_ott26_v1"
    if await db.app_meta.find_one({"key": key}):
        return
    path = seed_data.DATA_DIR / "clienti_menu_pellegrini_ott26.xlsx"
    if not path.exists():
        return
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    header = [seed_data._norm(h) if h is not None else "" for h in rows[0]]
    idx = {h: i for i, h in enumerate(header)}

    def cell(row, col):
        i = idx.get(col)
        v = row[i] if i is not None and i < len(row) else None
        if v is None:
            return ""
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return seed_data._norm(str(v))

    giri = await db.giri.find({"active": {"$ne": False}}, {"_id": 0}).to_list(1000)

    def resolve(agent, zona):
        k = seed_data._key(zona)
        if not k:
            return None, None
        mine = [g for g in giri if (g.get("agent") or "umberto") == agent]
        hits = []
        for g in mine:
            for pos, loc in enumerate(g.get("localities") or []):
                if seed_data._key(loc).replace("'", "") == k.replace("'", ""):
                    hits.append((g["id"], pos))
        if len({h[0] for h in hits}) == 1:
            return hits[0]
        if not hits:
            named = [g for g in mine if seed_data._key(g["name"]) == k]
            if len(named) == 1:
                return named[0]["id"], 999
        return None, None

    existing = await db.clients.find({"deleted_at": None}, {"_id": 0}).to_list(20000)
    by_piva = {_norm_piva((c.get("extra") or {}).get("partita_iva")): c for c in existing
               if _norm_piva((c.get("extra") or {}).get("partita_iva"))}
    by_code = {c.get("codice_azienda"): c for c in existing if c.get("codice_azienda")}
    by_name = {seed_data._key(c.get("ragione_sociale")): c for c in existing}

    inserted = updated = review = 0
    for row in rows[1:]:
        ragione = cell(row, "Ragione Sociale")
        if not ragione:
            continue
        agent = seed_data.AGENT_MAP.get(seed_data._key(cell(row, "Nominativo Agente")))
        if not agent:
            continue
        zona = cell(row, "Descrizione Zona")
        piva, cf = cell(row, "Partita IVA"), cell(row, "Codice Fiscale")
        tel = cell(row, "Telefono") or cell(row, "Telefono Cellulare") or cell(row, "Telefono Ufficio")
        giro_id, pos = resolve(agent, zona)
        cur = by_piva.get(_norm_piva(piva)) or by_code.get(cell(row, "Codice Azienda")) \
            or by_name.get(seed_data._key(ragione))
        if cur:
            upd = {"agent": agent}
            if piva:
                upd["extra.partita_iva"] = piva
            if cf:
                upd["extra.codice_fiscale"] = cf
            for f, v in (("indirizzo", cell(row, "Indirizzo")), ("cap", cell(row, "CAP")),
                         ("citta", cell(row, "Città")), ("provincia", cell(row, "Provincia")),
                         ("zona", zona), ("telefono", tel), ("email", cell(row, "Email")),
                         ("codice_azienda", cell(row, "Codice Azienda"))):
                if v and not cur.get(f):
                    upd[f] = v
            cur_giro = next((g for g in giri if g["id"] == cur.get("giro_id")), None)
            if not cur_giro or (cur_giro.get("agent") or "umberto") != agent:
                upd["giro_id"] = giro_id
                upd["position"] = pos if pos is not None else 999
                if not giro_id:
                    review += 1
            await db.clients.update_one({"id": cur["id"]}, {"$set": upd})
            updated += 1
            continue
        if not giro_id:
            review += 1
        await db.clients.insert_one({
            "id": str(uuid.uuid4()),
            "ragione_sociale": ragione,
            "codice_azienda": cell(row, "Codice Azienda"),
            "provincia": cell(row, "Provincia"),
            "giro_id": giro_id,
            "position": pos if pos is not None else 999,
            "citta": cell(row, "Città"),
            "zona": zona,
            "indirizzo": cell(row, "Indirizzo"),
            "cap": cell(row, "CAP"),
            "telefono": tel,
            "email": cell(row, "Email"),
            "agent": agent,
            "permanent_note": "",
            "last_visit_at": None,
            "snoozed_until": None,
            "extra": {
                "partita_iva": piva, "codice_fiscale": cf,
                "telefono_ufficio": cell(row, "Telefono Ufficio"),
                "telefono_cellulare": cell(row, "Telefono Cellulare"),
                "email_pec": cell(row, "Email PEC"),
                "sito_web": cell(row, "Sito Web"),
                "categoria": cell(row, "Descrizione Cat. Commerciali"),
                "pagamento": cell(row, "Descrizione Mod. di Pagamento"),
                "note_interne": cell(row, "Note Interne"),
                "codice_sdi": cell(row, "Codice SDI"),
            },
            "created_by": agent,
            "deleted_at": None,
            "created_at": now_utc(),
        })
        inserted += 1
    await db.app_meta.insert_one({"key": key, "at": now_utc(), "inserted": inserted,
                                  "updated": updated, "da_verificare": review})
    logger.info("Import clienti ott26: %d nuovi, %d aggiornati, %d da verificare", inserted, updated, review)


def _norm(s) -> str:
    if s is None:
        return ""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    s = s.lower().strip()
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


async def seed_recurrences():
    # Seed recurrence definitions (idempotent by company key).
    for d in recurrence_seed.RECURRENCE_DEFS:
        existing = await db.recurrence_defs.find_one({"company": d["company"]})
        if not existing:
            await db.recurrence_defs.insert_one({
                "id": str(uuid.uuid4()),
                "company": d["company"],
                "label": d["label"],
                "periods": d["periods"],
                "groups": d["groups"],
                "order": d["order"],
                "created_at": now_utc(),
            })

    # Import members once (guard marker).
    marker = await db.app_meta.find_one({"key": "recurrences_seeded_v1"})
    if marker:
        return

    # Build lookup of existing active clients by normalized name(+city).
    existing = await db.clients.find({"deleted_at": None}, {"_id": 0, "id": 1, "ragione_sociale": 1, "citta": 1, "agent": 1}).to_list(20000)
    by_namecity = {}
    for c in existing:
        key = _norm(c.get("ragione_sociale")) + "|" + _norm(c.get("citta"))
        by_namecity.setdefault(key, c)

    total, matched, created = 0, 0, 0
    for company, members in recurrence_seed.RECURRENCE_MEMBERS.items():
        pos_by_group = {}
        for m in members:
            total += 1
            key = _norm(m["ragione_sociale"]) + "|" + _norm(m["citta"])
            match = by_namecity.get(key)
            if match:
                client_id = match["id"]
                agent = match.get("agent") or m["agent"]
                matched += 1
            else:
                # Recurrence-only new client: no territorial giro, flagged for review
                # unless explicitly declared new in the prompt.
                client_id = str(uuid.uuid4())
                agent = m["agent"]
                await db.clients.insert_one({
                    "id": client_id,
                    "ragione_sociale": m["ragione_sociale"],
                    "codice_azienda": m.get("codice_azienda", ""),
                    "provincia": m.get("provincia", ""),
                    "giro_id": None,
                    "position": 999,
                    "citta": m.get("citta", ""),
                    "zona": m.get("citta", ""),
                    "indirizzo": m.get("indirizzo", ""),
                    "cap": m.get("cap", ""),
                    "telefono": m.get("telefono", ""),
                    "email": m.get("email", ""),
                    "agent": agent,
                    "permanent_note": "",
                    "last_visit_at": None,
                    "snoozed_until": None,
                    "extra": {
                        "recurrence_only": True,
                        "needs_review": not m.get("explicit_new", False),
                        "partita_iva": m.get("partita_iva", ""),
                        "codice_fiscale": m.get("codice_fiscale", ""),
                    },
                    "created_by": agent,
                    "deleted_at": None,
                    "created_at": now_utc(),
                })
                created += 1
            g = m["group"]
            pos = pos_by_group.get(g, 0)
            pos_by_group[g] = pos + 1
            # Avoid duplicate membership for same (company, client).
            dup = await db.recurrence_members.find_one({"company": company, "client_id": client_id})
            if not dup:
                await db.recurrence_members.insert_one({
                    "id": str(uuid.uuid4()),
                    "company": company,
                    "client_id": client_id,
                    "group": g,
                    "position": pos,
                    "agent": agent,
                    "created_at": now_utc(),
                })

    await db.app_meta.insert_one({"key": "recurrences_seeded_v1", "at": now_utc()})
    logger.info("Recurrences seeded: %d members (%d matched existing, %d created new)", total, matched, created)


@app.on_event("startup")
async def on_startup():
    # Su Vercel (serverless) lo startup gira a ogni cold start: NON eseguiamo il
    # seed pesante dei dati (il DB Atlas di produzione è già popolato).
    if os.environ.get("VERCEL") or os.environ.get("DISABLE_SEED"):
        return
    try:
        await seed()
    except Exception as e:
        logger.exception("SEED dati fallito: %s", e)


@app.on_event("shutdown")
async def on_shutdown():
    client.close()


# ---------------------------------------------------------------------------
# Auth routes
# ---------------------------------------------------------------------------
@api.get("/")
async def root():
    return {"message": "AgendaVisite API"}


@api.post("/auth/login")
async def login(body: LoginRequest):
    ident = body.username.strip().lower()
    # Accetta sia lo username (es. 'umberto') sia l'email (case-insensitive).
    user = await db.users.find_one({"$or": [{"username": ident}, {"email": ident}]})
    dummy = "$2b$12$" + "x" * 53
    if not user:
        verify_pw(body.password, dummy)
        logger.info("LOGIN FALLITO: identificativo '%s' non trovato", ident)
        raise HTTPException(status_code=401, detail="Username o password errati")
    if not user.get("is_active", False) or not verify_pw(body.password, user["hashed_password"]):
        logger.info("LOGIN FALLITO: password errata per '%s'", ident)
        raise HTTPException(status_code=401, detail="Username o password errati")
    token = create_token(user["username"])
    return {
        "access_token": token,
        "token_type": "bearer",
        "user": {"username": user["username"], "display_name": user["display_name"], "role": user.get("role", "agent")},
    }


@api.get("/auth/me")
async def me(user=Depends(get_current_user)):
    return {"username": user["username"], "display_name": user["display_name"], "role": user.get("role", "agent")}


async def require_admin(user=Depends(get_current_user)):
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Solo l'amministratore può eseguire questa operazione")
    return user


# ---------------------------------------------------------------------------
# Companies
# ---------------------------------------------------------------------------
# Ordine di priorità richiesto per l'elenco aziende/produttori.
COMPANY_PRIORITY = [
    "Librandi", "Serracavallo", "Pellegrini", "Villani", "Menù", "Cala",
    "Mazzetti d'Altavilla", "Toso", "Mosnel", "Tramin", "Pio Cesare",
    "Biondi Santi", "Piper-Heidsieck", "Bonfissuto", "Foss Marai", "Isole e Olena",
]
_COMPANY_PRIORITY_INDEX = {name: i for i, name in enumerate(COMPANY_PRIORITY)}


@api.get("/companies")
async def list_companies(include_inactive: bool = False, user=Depends(get_current_user)):
    q = {} if include_inactive else {"active": True}
    docs = await db.companies.find(q, {"_id": 0}).to_list(1000)
    docs.sort(key=lambda d: (
        _COMPANY_PRIORITY_INDEX.get(d.get("name"), len(COMPANY_PRIORITY)),
        d.get("name", "").lower(),
    ))
    return docs


@api.post("/companies")
async def create_company(body: Company, user=Depends(require_admin)):
    name = body.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Nome azienda obbligatorio")
    existing = await db.companies.find_one({"name": name})
    if existing:
        raise HTTPException(status_code=400, detail="Azienda già esistente")
    last = await db.companies.find_one({}, {"_id": 0}, sort=[("order", -1)])
    order = (last["order"] + 1) if last else 0
    doc = {"id": str(uuid.uuid4()), "name": name, "active": True, "order": order, "created_at": now_utc()}
    await db.companies.insert_one(doc)
    doc.pop("_id", None)
    doc.pop("created_at", None)
    return doc


@api.put("/companies/{company_id}")
async def update_company(company_id: str, body: CompanyUpdate, user=Depends(require_admin)):
    update = {k: v for k, v in body.dict().items() if v is not None}
    if not update:
        raise HTTPException(status_code=400, detail="Nessun dato")
    res = await db.companies.update_one({"id": company_id}, {"$set": update})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Azienda non trovata")
    doc = await db.companies.find_one({"id": company_id}, {"_id": 0})
    return doc


# ---------------------------------------------------------------------------
# Giri
# ---------------------------------------------------------------------------
@api.get("/giri")
async def list_giri(include_inactive: bool = False, user=Depends(get_current_user)):
    q = {} if include_inactive else {"active": True}
    # Giri per-agente: Andrea vede solo i propri; gli altri (Umberto/admin) vedono
    # i giri non appartenenti ad Andrea (inclusi quelli storici senza campo agent).
    if user["username"] == "andrea":
        q["agent"] = "andrea"
    else:
        q["agent"] = {"$ne": "andrea"}
    docs = await db.giri.find(q, {"_id": 0}).sort("order", 1).to_list(1000)
    # Conteggio clienti attivi assegnati a ciascun giro (per la UI: "N clienti").
    giro_ids = [g["id"] for g in docs]
    counts: dict = {}
    if giro_ids:
        pipeline = [
            {"$match": {"giro_id": {"$in": giro_ids}, "deleted_at": None}},
            {"$group": {"_id": "$giro_id", "n": {"$sum": 1}}},
        ]
        async for row in db.clients.aggregate(pipeline):
            counts[row["_id"]] = row["n"]
    for g in docs:
        g["client_count"] = counts.get(g["id"], 0)
    return docs


def _can_edit_giro(user, giro) -> bool:
    is_andrea_user = user["username"] == "andrea"
    is_andrea_giro = giro.get("agent") == "andrea"
    return is_andrea_user == is_andrea_giro


@api.post("/giri")
async def create_giro(body: GiroCreate, user=Depends(get_current_user)):
    name = body.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Nome giro obbligatorio")
    owner = "andrea" if user["username"] == "andrea" else user["username"]
    last = await db.giri.find_one({"agent": owner}, {"_id": 0}, sort=[("order", -1)])
    order = (last["order"] + 1) if last else 0
    doc = {"id": str(uuid.uuid4()), "name": name, "localities": body.localities,
           "order": order, "active": True, "agent": owner, "created_at": now_utc()}
    await db.giri.insert_one(doc)
    doc.pop("_id", None)
    doc.pop("created_at", None)
    return doc


@api.put("/giri/{giro_id}")
async def update_giro(giro_id: str, body: GiroUpdate, user=Depends(get_current_user)):
    giro = await db.giri.find_one({"id": giro_id}, {"_id": 0})
    if not giro:
        raise HTTPException(status_code=404, detail="Giro non trovato")
    if not _can_edit_giro(user, giro):
        raise HTTPException(status_code=403, detail="Non puoi modificare questo giro")
    update = {k: v for k, v in body.dict().items() if v is not None}
    if not update:
        raise HTTPException(status_code=400, detail="Nessun dato")
    await db.giri.update_one({"id": giro_id}, {"$set": update})
    doc = await db.giri.find_one({"id": giro_id}, {"_id": 0})
    return doc


@api.delete("/giri/{giro_id}")
async def delete_giro(giro_id: str, user=Depends(get_current_user)):
    giro = await db.giri.find_one({"id": giro_id}, {"_id": 0})
    if not giro:
        raise HTTPException(status_code=404, detail="Giro non trovato")
    if not _can_edit_giro(user, giro):
        raise HTTPException(status_code=403, detail="Non puoi eliminare questo giro")
    # Sicurezza: non eliminare se contiene ancora clienti (vanno prima riassegnati).
    n = await db.clients.count_documents({"giro_id": giro_id, "agent": user["username"], "deleted_at": None})
    if n > 0:
        raise HTTPException(status_code=400, detail=f"Il giro contiene {n} clienti: spostali prima di eliminarlo")
    await db.giri.update_one({"id": giro_id}, {"$set": {"active": False}})
    return {"ok": True}


# ---------------------------------------------------------------------------
# Clients
# ---------------------------------------------------------------------------
async def _handled_month_ids(client_ids: List[str]) -> set:
    """Clienti con almeno un'attività (ordine, visita, incasso, nota) nel MESE
    corrente. Base per il conteggio GESTITI mensile (reset il 1° del mese)."""
    if not client_ids:
        return set()
    start = start_of_month_utc()
    cursor = db.events.find(
        {"client_id": {"$in": client_ids}, "created_at": {"$gte": start}, "deleted_at": None,
         "type": {"$in": ["order", "visit", "collection", "note"]}},
        {"_id": 0, "client_id": 1},
    )
    ids = set()
    async for e in cursor:
        ids.add(e["client_id"])
    return ids


async def _active_suspensions_for(client_ids: List[str]) -> dict:
    """Return {client_id: [company_name,...]} of currently active suspensions.
    A suspension (manual event or overdue deferred order) for a company is active
    unless a later 'collection' (incassato) for that same company settled it.
    """
    if not client_ids:
        return {}
    now = now_utc()
    events = await db.events.find(
        {"client_id": {"$in": client_ids}, "deleted_at": None,
         "type": {"$in": ["order", "suspension", "collection"]}},
        {"_id": 0},
    ).to_list(200000)

    from collections import defaultdict

    def _aware(dt):
        if dt is None:
            return None
        return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)

    data = defaultdict(lambda: defaultdict(lambda: {"paid": None, "manual": [], "auto": []}))
    for e in events:
        comp = e.get("company_name") or e.get("company_id")
        if not comp:
            continue
        cid = e["client_id"]
        created = _aware(e.get("created_at"))
        slot = data[cid][comp]
        if e["type"] == "collection":
            if slot["paid"] is None or (created and created > slot["paid"]):
                slot["paid"] = created
        elif e["type"] == "suspension":
            slot["manual"].append(created)
        elif e["type"] == "order":
            due = _aware(e.get("due_at"))
            if due is not None and due <= now:
                slot["auto"].append(created)

    result = {}
    for cid, comps in data.items():
        active = []
        for comp, slot in comps.items():
            paid = slot["paid"]
            man = any(paid is None or (c and c > paid) for c in slot["manual"])
            auto = any(paid is None or (c and c > paid) for c in slot["auto"])
            if man or auto:
                active.append(comp)
        if active:
            result[cid] = sorted(active, key=lambda x: x.lower())
    return result


async def _last_events_for(client_ids: List[str], types: List[str]) -> dict:
    """Return {type: {client_id: created_at(datetime)}} of the most recent event
    of each given type. Persistent across months / 21-day cycle resets."""
    out = {t: {} for t in types}
    if not client_ids:
        return out
    events = await db.events.find(
        {"client_id": {"$in": client_ids}, "deleted_at": None, "type": {"$in": types}},
        {"_id": 0, "client_id": 1, "created_at": 1, "type": 1},
    ).to_list(200000)
    for e in events:
        t = e["type"]
        cid = e["client_id"]
        cur = e.get("created_at")
        if cur is None:
            continue
        d = out[t]
        if cid not in d or cur > d[cid]:
            d[cid] = cur
    return out


def _compute_status(doc: dict, handled_month: set) -> str:
    """Logica MENSILE: un cliente è GESTITO se ha avuto almeno un'attività
    (ordine/visita/incasso/nota) nel mese corrente, oppure se è stato rimandato
    a una data futura. Altrimenti è DA VISITARE. Reset automatico il 1° del mese."""
    now = now_utc()
    if doc["id"] in handled_month:
        return "gestito"
    snoozed = doc.get("snoozed_until")
    if snoozed is not None:
        if snoozed.tzinfo is None:
            snoozed = snoozed.replace(tzinfo=timezone.utc)
        if snoozed > now:
            return "gestito"
    return "da_visitare"


@api.get("/clients")
async def list_clients(giro_id: str = Query(...), user=Depends(get_current_user)):
    # Prendiamo TUTTI i clienti dell'agente: una sede (principale o aggiuntiva)
    # può appartenere a questo giro anche se la sede principale è altrove.
    docs = await db.clients.find(
        {"agent": user["username"], "deleted_at": None}, {"_id": 0}
    ).to_list(20000)
    ids = [d["id"] for d in docs]
    handled = await _handled_month_ids(ids)
    suspensions = await _active_suspensions_for(ids)
    last_ev = await _last_events_for(ids, ["order", "collection"])

    rows = []
    for d in docs:
        base = client_public(d)
        base["status"] = _compute_status(d, handled)
        base["handled_this_month"] = d["id"] in handled
        base["suspensions"] = suspensions.get(d["id"], [])
        base["last_order_at"] = iso(last_ev["order"].get(d["id"]))
        base["last_collection_at"] = iso(last_ev["collection"].get(d["id"]))

        # Sede principale (dati anagrafici del cliente)
        if d.get("giro_id") == giro_id:
            r = dict(base)
            r["row_id"] = f'{d["id"]}:principal'
            r["site_id"] = "principal"
            r["site_nominativo"] = (d.get("extra") or {}).get("sede_nome") or ""
            r["position"] = d.get("position")
            rows.append(r)

        # Sedi di scarico aggiuntive assegnate a questo giro
        for s in (d.get("extra") or {}).get("sites", []):
            if s.get("giro_id") == giro_id:
                r = dict(base)
                r["row_id"] = f'{d["id"]}:{s.get("id")}'
                r["site_id"] = s.get("id")
                r["site_nominativo"] = s.get("nominativo") or ""
                r["indirizzo"] = s.get("indirizzo") or ""
                r["citta"] = s.get("citta") or ""
                r["zona"] = s.get("zona") or ""
                r["cap"] = s.get("cap") or ""
                r["provincia"] = s.get("provincia") or ""
                r["position"] = s.get("position")
                rows.append(r)

    rows.sort(key=lambda r: (
        r.get("position") if r.get("position") is not None else 999,
        r.get("ragione_sociale", "").lower(),
    ))
    return rows


@api.get("/clients/da-verificare")
async def da_verificare(user=Depends(get_current_user)):
    q = {"deleted_at": None, "$or": [{"giro_id": None}, {"extra.needs_review": True}]}
    if user.get("role") != "admin":
        q["agent"] = user["username"]
    docs = await db.clients.find(q, {"_id": 0}).to_list(5000)
    docs.sort(key=lambda d: d.get("ragione_sociale", "").lower())

    # Possibili duplicati: altri clienti ATTIVI con la stessa Partita IVA (normalizzata).
    def _np(s):
        return re.sub(r"\s+", "", (s or "").strip().upper())

    from collections import defaultdict
    allc = await db.clients.find(
        {"deleted_at": None}, {"_id": 0, "id": 1, "ragione_sociale": 1, "citta": 1, "agent": 1, "extra": 1}
    ).to_list(20000)
    by_piva = defaultdict(list)
    for c in allc:
        p = _np((c.get("extra") or {}).get("partita_iva"))
        if p and len(p) >= 5:
            by_piva[p].append(c)

    out = []
    for d in docs:
        pub = client_public(d)
        p = _np((d.get("extra") or {}).get("partita_iva"))
        dups = []
        if p and len(p) >= 5:
            dups = [
                {"id": c["id"], "ragione_sociale": c.get("ragione_sociale", ""),
                 "citta": c.get("citta", ""), "agent": c.get("agent", "")}
                for c in by_piva.get(p, []) if c["id"] != d["id"]
            ]
        pub["partita_iva"] = (d.get("extra") or {}).get("partita_iva", "")
        pub["duplicates"] = dups
        out.append(pub)
    return out


VALID_AGENTS = {"umberto", "andrea"}


@api.get("/clients/all")
async def list_all_clients(search: str = Query("", alias="search"), user=Depends(get_current_user)):
    q = {"deleted_at": None}
    if user.get("role") != "admin":
        q["agent"] = user["username"]
    s = (search or "").strip()
    if s:
        q["ragione_sociale"] = {"$regex": re.escape(s), "$options": "i"}
    docs = await db.clients.find(q, {"_id": 0}).to_list(10000)
    docs.sort(key=lambda d: d.get("ragione_sociale", "").lower())
    return [client_public(d) for d in docs]


@api.get("/clients/search")
async def search_clients(q: str = Query("", alias="q"), user=Depends(get_current_user)):
    """Ricerca libera (nome O comune) su TUTTI i clienti dell'agente, senza
    bisogno di selezionare un giro. Risultati arricchiti con stato/sospesi/
    ultime azioni come la lista del giro. Limite 60 risultati."""
    s = (q or "").strip()
    if not s:
        return []
    query: dict = {"deleted_at": None}
    if user.get("role") != "admin":
        query["agent"] = user["username"]
    rx = {"$regex": re.escape(s), "$options": "i"}
    query["$or"] = [{"ragione_sociale": rx}, {"citta": rx}]
    docs = await db.clients.find(query, {"_id": 0}).to_list(400)
    docs.sort(key=lambda d: d.get("ragione_sociale", "").lower())
    docs = docs[:60]
    ids = [d["id"] for d in docs]
    handled = await _handled_month_ids(ids)
    suspensions = await _active_suspensions_for(ids)
    last_ev = await _last_events_for(ids, ["order", "collection"])
    result = []
    for d in docs:
        pub = client_public(d)
        pub["status"] = _compute_status(d, handled)
        pub["handled_this_month"] = d["id"] in handled
        pub["suspensions"] = suspensions.get(d["id"], [])
        pub["last_order_at"] = iso(last_ev["order"].get(d["id"]))
        pub["last_collection_at"] = iso(last_ev["collection"].get(d["id"]))
        result.append(pub)
    return result


@api.post("/clients")
async def create_client(body: ClientCreate, user=Depends(get_current_user)):
    if not body.ragione_sociale.strip():
        raise HTTPException(status_code=400, detail="Ragione sociale obbligatoria")
    agent = user["username"]
    if user.get("role") == "admin" and body.agent in VALID_AGENTS:
        agent = body.agent
    doc = {
        "id": str(uuid.uuid4()),
        "ragione_sociale": body.ragione_sociale.strip(),
        "codice_azienda": "",
        "provincia": normalize_provincia(body.provincia),
        "giro_id": body.giro_id,
        "position": body.position if body.position is not None else 999,
        "citta": body.citta.strip(),
        "zona": body.zona.strip(),
        "indirizzo": body.indirizzo.strip(),
        "cap": body.cap.strip(),
        "telefono": body.telefono.strip(),
        "email": body.email.strip(),
        "agent": agent,
        "permanent_note": "",
        "last_visit_at": None,
        "snoozed_until": None,
        "extra": {
            "partita_iva": (body.partita_iva or "").strip(),
            "sede_nome": (body.sede_nome or "").strip(),
            "goose_code": (body.goose_code or "").strip(),
            "sites": normalize_sites(body.sites),
        },
        "created_by": user["username"],
        "deleted_at": None,
        "created_at": now_utc(),
    }
    await db.clients.insert_one(doc)
    return client_public(doc)


@api.get("/clients/{client_id}")
async def get_client(client_id: str, user=Depends(get_current_user)):
    doc = await db.clients.find_one({"id": client_id, "deleted_at": None}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    pub = client_public(doc)
    susp = await _active_suspensions_for([client_id])
    pub["suspensions"] = susp.get(client_id, [])
    return pub


@api.put("/clients/{client_id}")
async def update_client(client_id: str, body: ClientUpdate, user=Depends(get_current_user)):
    update = {k: v for k, v in body.dict(exclude_unset=True).items()}
    if not update:
        raise HTTPException(status_code=400, detail="Nessun dato")
    if "provincia" in update:
        update["provincia"] = normalize_provincia(update["provincia"])
    target = await db.clients.find_one({"id": client_id, "deleted_at": None}, {"_id": 0})
    if not target:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    if user.get("role") != "admin" and target.get("agent") != user["username"]:
        raise HTTPException(status_code=403, detail="Non puoi modificare questo cliente")
    # Only admin may (re)assign the agent; validate value.
    if "agent" in update:
        if user.get("role") != "admin" or update["agent"] not in VALID_AGENTS:
            update.pop("agent", None)
    set_doc = dict(update)
    if "partita_iva" in set_doc:
        set_doc["extra.partita_iva"] = (set_doc.pop("partita_iva") or "").strip()
    if "sede_nome" in set_doc:
        set_doc["extra.sede_nome"] = (set_doc.pop("sede_nome") or "").strip()
    if "goose_code" in set_doc:
        set_doc["extra.goose_code"] = (set_doc.pop("goose_code") or "").strip()
    if "sites" in set_doc:
        set_doc.pop("sites", None)
        set_doc["extra.sites"] = normalize_sites(update.get("sites"))
    if update.get("giro_id"):
        set_doc["extra.needs_review"] = False
    await db.clients.update_one({"id": client_id, "deleted_at": None}, {"$set": set_doc})
    doc = await db.clients.find_one({"id": client_id}, {"_id": 0})
    return client_public(doc)


@api.delete("/clients/{client_id}")
async def delete_client(client_id: str, user=Depends(get_current_user)):
    target = await db.clients.find_one({"id": client_id, "deleted_at": None})
    if not target:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    # L'admin può eliminare qualsiasi cliente; l'agente solo i propri.
    if user.get("role") != "admin" and target.get("agent") != user["username"]:
        raise HTTPException(status_code=403, detail="Non autorizzato")
    await db.clients.update_one(
        {"id": client_id}, {"$set": {"deleted_at": now_utc()}}
    )
    return {"ok": True}


class MergeClientBody(BaseModel):
    target_id: str


async def _merge_client_into(source_id: str, target_id: str):
    """Riassegna eventi e ricorrenze dal cliente sorgente al target, completa i
    campi mancanti del target e azzera needs_review. NON cancella il sorgente."""
    await db.events.update_many(
        {"client_id": source_id, "deleted_at": None}, {"$set": {"client_id": target_id}}
    )
    async for m in db.recurrence_members.find({"client_id": source_id}):
        exists = await db.recurrence_members.find_one(
            {"client_id": target_id, "company": m.get("company"), "period": m.get("period")}
        )
        if exists:
            await db.recurrence_members.delete_one({"_id": m["_id"]})
        else:
            await db.recurrence_members.update_one({"_id": m["_id"]}, {"$set": {"client_id": target_id}})
    src = await db.clients.find_one({"id": source_id}, {"_id": 0}) or {}
    tgt = await db.clients.find_one({"id": target_id}, {"_id": 0}) or {}
    fill = {}
    for f in ["telefono", "email", "indirizzo", "cap", "citta", "zona", "provincia", "permanent_note"]:
        if not (tgt.get(f) or "").strip() and (src.get(f) or "").strip():
            fill[f] = src[f]
    textra = dict(tgt.get("extra") or {})
    for k, v in (src.get("extra") or {}).items():
        if not (textra.get(k) or "") and v:
            textra[k] = v
    textra["needs_review"] = False
    fill["extra"] = textra
    await db.clients.update_one({"id": target_id}, {"$set": fill})


@api.post("/clients/{client_id}/merge")
async def merge_client(client_id: str, body: MergeClientBody, user=Depends(require_admin)):
    """Fonde il cliente {client_id} (sorgente) nel cliente {target_id}: lo storico
    passa al target, il sorgente viene soft-deleted. Solo admin."""
    if client_id == body.target_id:
        raise HTTPException(status_code=400, detail="Sorgente e destinazione coincidono")
    src = await db.clients.find_one({"id": client_id, "deleted_at": None}, {"_id": 0})
    tgt = await db.clients.find_one({"id": body.target_id, "deleted_at": None}, {"_id": 0})
    if not src or not tgt:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    await _merge_client_into(client_id, body.target_id)
    await db.clients.update_one({"id": client_id}, {"$set": {"deleted_at": now_utc()}})
    doc = await db.clients.find_one({"id": body.target_id}, {"_id": 0})
    return client_public(doc)


@api.get("/clients/{client_id}/history")
async def client_history(client_id: str, user=Depends(get_current_user)):
    docs = await db.events.find(
        {"client_id": client_id, "deleted_at": None}, {"_id": 0}
    ).sort("created_at", -1).to_list(5000)
    return [event_public(d) for d in docs]


# ---------------------------------------------------------------------------
# Events (quick actions)
# ---------------------------------------------------------------------------
VALID_EVENT_TYPES = {"visit", "order", "reschedule", "collection", "note", "suspension"}


@api.get("/payment-modes")
async def payment_modes(user=Depends(get_current_user)):
    return [{"key": k, "label": PAYMENT_MODES[k]["label"], "days": PAYMENT_MODES[k]["days"]} for k in PAYMENT_MODE_ORDER]


@api.post("/events")
async def create_event(body: EventCreate, user=Depends(get_current_user)):
    if body.type not in VALID_EVENT_TYPES:
        raise HTTPException(status_code=400, detail="Tipo evento non valido")
    cli = await db.clients.find_one({"id": body.client_id, "deleted_at": None}, {"_id": 0})
    if not cli:
        raise HTTPException(status_code=404, detail="Cliente non trovato")

    now = now_utc()
    # Data dell'attività scelta dall'utente (default: adesso). Consente di
    # registrare ordini/visite/incassi effettuati nei giorni precedenti.
    base = parse_activity_date(body.activity_date, now)
    doc = {
        "id": str(uuid.uuid4()),
        "client_id": body.client_id,
        "type": body.type,
        "company_id": None,
        "company_name": None,
        "note_text": None,
        "reschedule_until": None,
        "payment_mode": None,
        "due_at": None,
        "source": None,
        "collection_method": None,
        "collection_ref_date": None,
        "site_id": body.site_id,
        "site_nominativo": body.site_nominativo,
        "site_indirizzo": body.site_indirizzo,
        "agent": user["username"],
        "created_at": base,
        "deleted_at": None,
    }

    if body.type in ("order", "collection", "suspension"):
        if not body.company_id:
            raise HTTPException(status_code=400, detail="Seleziona un'azienda")
        comp = await db.companies.find_one({"id": body.company_id}, {"_id": 0})
        if not comp:
            raise HTTPException(status_code=404, detail="Azienda non trovata")
        doc["company_id"] = comp["id"]
        doc["company_name"] = comp["name"]

    if body.type == "order":
        # Payment mode chosen for THIS order only (does not touch client anagrafica).
        if body.payment_mode:
            if body.payment_mode not in PAYMENT_MODES:
                raise HTTPException(status_code=400, detail="Modalità di pagamento non valida")
            doc["payment_mode"] = body.payment_mode
            days = PAYMENT_MODES[body.payment_mode]["days"]
            if days is not None:
                doc["due_at"] = base + timedelta(days=days)

    if body.type == "suspension":
        doc["source"] = "manual"

    if body.type == "collection":
        # L'incasso salda tutto ciò che risulta in sospeso FINO al giorno scelto
        # (incluso): usiamo fine giornata (ora di Roma) come timestamp di confronto.
        day_end = base.astimezone(ROME).replace(hour=23, minute=59, second=59, microsecond=0)
        doc["created_at"] = day_end.astimezone(timezone.utc)
        if body.collection_method:
            if body.collection_method not in ("contanti", "bonifico", "assegno"):
                raise HTTPException(status_code=400, detail="Modalità di incasso non valida")
            doc["collection_method"] = body.collection_method
            if body.collection_method == "bonifico" and body.collection_ref_date:
                doc["collection_ref_date"] = body.collection_ref_date

    if body.type == "note":
        if not (body.note_text and body.note_text.strip()):
            raise HTTPException(status_code=400, detail="Nota vuota")
        doc["note_text"] = body.note_text.strip()

    if body.type == "visit" and body.note_text and body.note_text.strip():
        doc["note_text"] = body.note_text.strip()[:200]
    if body.type == "reschedule":
        target = None
        if body.reschedule_date:
            try:
                target = datetime.fromisoformat(body.reschedule_date.replace("Z", "+00:00"))
                if target.tzinfo is None:
                    target = target.replace(tzinfo=timezone.utc)
            except ValueError:
                raise HTTPException(status_code=400, detail="Data non valida")
        elif body.reschedule_days:
            target = now + timedelta(days=body.reschedule_days)
        doc["reschedule_until"] = target
        if target is not None:
            await db.clients.update_one({"id": body.client_id}, {"$set": {"snoozed_until": target}})

    await db.events.insert_one(doc)

    if body.type == "visit":
        await db.clients.update_one(
            {"id": body.client_id},
            {"$set": {"last_visit_at": base, "snoozed_until": None}},
        )

    return event_public(doc)


class EventUpdate(BaseModel):
    payment_mode: Optional[str] = None
    activity_date: Optional[str] = None
    collection_method: Optional[str] = None
    collection_ref_date: Optional[str] = None


async def _event_auth(event_id: str, user) -> dict:
    ev = await db.events.find_one({"id": event_id, "deleted_at": None}, {"_id": 0})
    if not ev:
        raise HTTPException(status_code=404, detail="Evento non trovato")
    cli = await db.clients.find_one({"id": ev["client_id"]}, {"_id": 0})
    if user.get("role") != "admin" and (not cli or cli.get("agent") != user["username"]):
        raise HTTPException(status_code=403, detail="Non autorizzato")
    return ev


@api.put("/events/{event_id}")
async def update_event(event_id: str, body: EventUpdate, user=Depends(get_current_user)):
    """Modifica un evento già registrato. Per gli ordini: termini di pagamento
    (ricalcola scadenza) e/o data. Per gli incassi: modalità e/o data."""
    ev = await _event_auth(event_id, user)
    update: dict = {}
    created = ev.get("created_at")
    if created is not None and created.tzinfo is None:
        created = created.replace(tzinfo=timezone.utc)
    if body.activity_date:
        created = parse_activity_date(body.activity_date, created or now_utc())
        if ev["type"] == "collection":
            # Come alla creazione: fine giornata (Roma) per saldare fino a quel giorno.
            created = created.astimezone(ROME).replace(hour=23, minute=59, second=59, microsecond=0).astimezone(timezone.utc)
        update["created_at"] = created
    if ev["type"] == "order":
        pm = body.payment_mode if body.payment_mode is not None else ev.get("payment_mode")
        if pm is not None and pm not in PAYMENT_MODES:
            raise HTTPException(status_code=400, detail="Modalità di pagamento non valida")
        if body.payment_mode is not None:
            update["payment_mode"] = pm
        if ("payment_mode" in update) or ("created_at" in update):
            days = PAYMENT_MODES.get(pm, {}).get("days") if pm else None
            update["due_at"] = (created + timedelta(days=days)) if (days is not None and created) else None
    if ev["type"] == "collection" and body.collection_method is not None:
        if body.collection_method not in ("contanti", "bonifico", "assegno"):
            raise HTTPException(status_code=400, detail="Modalità di incasso non valida")
        update["collection_method"] = body.collection_method
        update["collection_ref_date"] = body.collection_ref_date if body.collection_method == "bonifico" else None
    if not update:
        raise HTTPException(status_code=400, detail="Nessuna modifica")
    await db.events.update_one({"id": event_id}, {"$set": update})
    doc = await db.events.find_one({"id": event_id}, {"_id": 0})
    return event_public(doc)


@api.delete("/events/{event_id}")
async def delete_event(event_id: str, user=Depends(get_current_user)):
    """Elimina (soft-delete) un evento: lo rimuove da storico, promemoria sospesi
    e banner informativo."""
    await _event_auth(event_id, user)
    await db.events.update_one({"id": event_id}, {"$set": {"deleted_at": now_utc()}})
    return {"ok": True}
ACTIVITY_TYPES = ["order", "collection", "suspension", "reschedule"]
ACTIVITY_TYPE_LABELS = {
    "order": "Ordine effettuato",
    "collection": "Incassato",
    "suspension": "+Sospeso",
    "reschedule": "Visita rimandata",
}


@api.get("/activities")
async def list_activities(
    scope: str = Query("all"),  # all | umberto | andrea
    type: str = Query("all"),   # all | order | collection | suspension | reschedule
    limit: int = Query(200),
    from_date: Optional[str] = Query(None),  # 'YYYY-MM-DD' (giorno locale Rome)
    to_date: Optional[str] = Query(None),    # 'YYYY-MM-DD' (incluso)
    user=Depends(get_current_user),
):
    is_admin = user.get("role") == "admin"
    q = {"deleted_at": None, "type": {"$in": ACTIVITY_TYPES}}
    if type != "all" and type in ACTIVITY_TYPES:
        q["type"] = type
    # Filtro periodo (inizio giornata <= created_at < inizio giorno successivo a to_date).
    dfrom = parse_day_bound(from_date, end=False)
    dto = parse_day_bound(to_date, end=True)
    if dfrom or dto:
        rng = {}
        if dfrom:
            rng["$gte"] = dfrom
        if dto:
            rng["$lt"] = dto
        q["created_at"] = rng
    # Permissions: agents only ever see their own activities regardless of scope.
    if not is_admin:
        q["agent"] = user["username"]
    elif scope in VALID_AGENTS:
        q["agent"] = scope
    lim = max(1, min(limit, 2000))
    events = await db.events.find(q, {"_id": 0}).sort("created_at", -1).limit(lim).to_list(lim)

    client_ids = list({e["client_id"] for e in events})
    clients = await db.clients.find({"id": {"$in": client_ids}}, {"_id": 0}).to_list(20000)
    cmap = {c["id"]: c for c in clients}
    giri = await db.giri.find({}, {"_id": 0, "id": 1, "name": 1}).to_list(1000)
    gmap = {g["id"]: g["name"] for g in giri}

    out = []
    for e in events:
        cli = cmap.get(e["client_id"], {})
        giro_name = gmap.get(cli.get("giro_id")) if cli else None
        type_label = ACTIVITY_TYPE_LABELS.get(e["type"], e["type"])
        if e["type"] == "order" and e.get("company_name"):
            type_label = f"Ordine effettuato per {e['company_name']}"
            context = f"Ordine N. {e.get('order_number')}" if e.get("order_number") else ""
        elif e["type"] in ("order", "collection", "suspension"):
            if e.get("order_number"):
                context = f"Ordine N. {e.get('order_number')}"
            else:
                context = e.get("company_name") or ""
        else:
            context = giro_name or (cli.get("citta") if cli else "") or ""
        site_label = e.get("ship_label") or e.get("site_nominativo") or None
        out.append({
            "id": e["id"],
            "type": e["type"],
            "type_label": type_label,
            "created_at": iso(e.get("created_at")),
            "agent": e.get("agent"),
            "client_id": e.get("client_id"),
            "client_ragione_sociale": cli.get("ragione_sociale", "") if cli else "",
            "context": context,
            "giro_name": giro_name,
            "citta": cli.get("citta", "") if cli else "",
            "company_name": e.get("company_name"),
            "payment_mode": e.get("payment_mode"),
            "site_nominativo": e.get("site_nominativo"),
            "order_number": e.get("order_number"),
            "site_label": site_label,
        })
    return out


@api.get("/suspensions")
async def list_suspensions(scope: str = Query("all"), user=Depends(get_current_user)):
    """Promemoria: clienti con sospesi attivi (scaduti/manuali non incassati) e
    incassi differiti IN SCADENZA entro DUE_SOON_DAYS. Un elemento per (cliente, azienda)."""
    is_admin = user.get("role") == "admin"
    cq = {"deleted_at": None}
    if not is_admin:
        cq["agent"] = user["username"]
    elif scope in VALID_AGENTS:
        cq["agent"] = scope
    clients = await db.clients.find(cq, {"_id": 0}).to_list(20000)
    cmap = {c["id"]: c for c in clients}
    cids = list(cmap.keys())
    giri = await db.giri.find({}, {"_id": 0, "id": 1, "name": 1}).to_list(1000)
    gmap = {g["id"]: g["name"] for g in giri}

    now = now_utc()
    soon = now + timedelta(days=DUE_SOON_DAYS)

    def _aware(dt):
        if dt is None:
            return None
        return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)

    from collections import defaultdict
    slots = defaultdict(lambda: {"paid": None, "manual": [], "orders": [], "company_id": None})
    if cids:
        events = await db.events.find(
            {"client_id": {"$in": cids}, "deleted_at": None,
             "type": {"$in": ["order", "suspension", "collection"]}},
            {"_id": 0},
        ).to_list(200000)
        for e in events:
            comp = e.get("company_name") or e.get("company_id")
            if not comp:
                continue
            s = slots[(e["client_id"], comp)]
            if e.get("company_id"):
                s["company_id"] = e["company_id"]
            created = _aware(e.get("created_at"))
            if e["type"] == "collection":
                if s["paid"] is None or (created and created > s["paid"]):
                    s["paid"] = created
            elif e["type"] == "suspension":
                s["manual"].append((created, e["id"]))
            elif e["type"] == "order":
                due = _aware(e.get("due_at"))
                if due is not None:
                    s["orders"].append((created, due, e["id"], e.get("payment_mode")))

    items = []
    for (cid, comp), s in slots.items():
        paid = s["paid"]
        open_manual = sorted(
            [(c, eid) for (c, eid) in s["manual"] if (paid is None or (c and c > paid))],
            key=lambda m: m[0] or now,
        )
        open_orders = [(c, due, eid, pm) for (c, due, eid, pm) in s["orders"] if (paid is None or (c and c > paid))]
        overdue_orders = sorted([o for o in open_orders if o[1] <= now], key=lambda o: o[1])
        duesoon_orders = sorted([o for o in open_orders if now < o[1] <= soon], key=lambda o: o[1])
        manual_since = open_manual[0][0] if open_manual else None

        cli = cmap.get(cid, {})
        base = {
            "client_id": cid,
            "ragione_sociale": cli.get("ragione_sociale", ""),
            "citta": cli.get("citta", ""),
            "agent": cli.get("agent", ""),
            "giro_name": gmap.get(cli.get("giro_id")),
            "company_name": comp,
            "company_id": s.get("company_id"),
        }
        if open_manual or overdue_orders:
            if overdue_orders:
                _c, due0, eid0, pm0 = overdue_orders[0]
                items.append({**base, "kind": "overdue", "due_at": iso(due0), "since": iso(manual_since),
                              "event_id": eid0, "event_type": "order", "payment_mode": pm0})
            else:
                items.append({**base, "kind": "overdue", "due_at": None, "since": iso(manual_since),
                              "event_id": open_manual[0][1], "event_type": "suspension", "payment_mode": None})
        elif duesoon_orders:
            _c, due0, eid0, pm0 = duesoon_orders[0]
            items.append({**base, "kind": "due_soon", "due_at": iso(due0), "since": None,
                          "event_id": eid0, "event_type": "order", "payment_mode": pm0})

    # Ordinamento: prima gli scaduti, poi in scadenza; per data di scadenza crescente.
    def _key(it):
        kind_rank = 0 if it["kind"] == "overdue" else 1
        return (kind_rank, it["due_at"] or "")
    items.sort(key=_key)
    return items


@api.get("/clients/{client_id}/pending-suspensions")
async def client_pending_suspensions(client_id: str, user=Depends(get_current_user)):
    """TUTTE le forniture in sospeso (non incassate) del cliente, indipendentemente
    dalla finestra di scadenza. Usato per il banner di promemoria quando si apre
    la scheda azioni di un cliente. Un elemento per (azienda)."""
    cli = await db.clients.find_one({"id": client_id, "deleted_at": None}, {"_id": 0})
    if not cli:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    if user.get("role") != "admin" and cli.get("agent") != user["username"]:
        raise HTTPException(status_code=403, detail="Non autorizzato")

    now = now_utc()

    def _aware(dt):
        if dt is None:
            return None
        return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)

    from collections import defaultdict
    slots = defaultdict(lambda: {"paid": None, "manual": [], "orders": [], "company_id": None})
    events = await db.events.find(
        {"client_id": client_id, "deleted_at": None,
         "type": {"$in": ["order", "suspension", "collection"]}},
        {"_id": 0},
    ).to_list(100000)
    for e in events:
        comp = e.get("company_name") or e.get("company_id")
        if not comp:
            continue
        s = slots[comp]
        if e.get("company_id"):
            s["company_id"] = e["company_id"]
        created = _aware(e.get("created_at"))
        if e["type"] == "collection":
            if s["paid"] is None or (created and created > s["paid"]):
                s["paid"] = created
        elif e["type"] == "suspension":
            s["manual"].append(created)
        elif e["type"] == "order":
            due = _aware(e.get("due_at"))
            if due is not None:
                s["orders"].append((created, due))

    items = []
    for comp, s in slots.items():
        paid = s["paid"]
        manual_active = any(paid is None or (c and c > paid) for c in s["manual"])
        manual_since = min((c for c in s["manual"] if c and (paid is None or c > paid)), default=None)
        open_dues = [due for created, due in s["orders"] if (paid is None or (created and created > paid))]
        if not manual_active and not open_dues:
            continue
        due_at = min(open_dues) if open_dues else None
        is_overdue = bool((due_at and due_at <= now) or (manual_active and due_at is None))
        items.append({
            "company_name": comp,
            "company_id": s.get("company_id"),
            "kind": "overdue" if is_overdue else "pending",
            "due_at": iso(due_at),
            "since": iso(manual_since),
        })
    items.sort(key=lambda it: (0 if it["kind"] == "overdue" else 1, it["due_at"] or "z"))
    return items


# ---------------------------------------------------------------------------
# Recurrences (Ricorrenze) — separate from territorial giri
# ---------------------------------------------------------------------------
class RecurrenceOrderBody(BaseModel):
    client_id: str
    period: str


class RecurrenceMemberCreate(BaseModel):
    group: str
    period: Optional[str] = None
    client_id: Optional[str] = None
    ragione_sociale: Optional[str] = None
    citta: Optional[str] = ""
    provincia: Optional[str] = ""
    indirizzo: Optional[str] = ""
    cap: Optional[str] = ""
    telefono: Optional[str] = ""
    email: Optional[str] = ""
    agent: Optional[str] = None


def _rome_year_start() -> datetime:
    now_rome = now_utc().astimezone(ROME)
    return datetime(now_rome.year, 1, 1, tzinfo=ROME).astimezone(timezone.utc)


async def _recurrence_done_ids(company: str, period: str, client_ids: List[str]) -> dict:
    """Return {client_id: order_iso_date} for orders placed this calendar year."""
    if not client_ids:
        return {}
    start = _rome_year_start()
    cursor = db.events.find(
        {"type": "recurrence_order", "recurrence_company": company,
         "recurrence_period": period, "client_id": {"$in": client_ids},
         "created_at": {"$gte": start}, "deleted_at": None},
        {"_id": 0, "client_id": 1, "created_at": 1},
    )
    out = {}
    async for e in cursor:
        cid = e["client_id"]
        if cid not in out:
            out[cid] = iso(e["created_at"])
    return out


async def _get_recur_def(company: str) -> dict:
    d = await db.recurrence_defs.find_one({"company": company}, {"_id": 0})
    if not d:
        raise HTTPException(status_code=404, detail="Ricorrenza non trovata")
    return d


@api.get("/recurrences")
async def list_recurrences(user=Depends(get_current_user)):
    defs = await db.recurrence_defs.find({}, {"_id": 0}).sort("order", 1).to_list(100)
    return defs


@api.get("/recurrences/{company}/members")
async def recurrence_members(company: str, period: str = Query(...), user=Depends(get_current_user)):
    rdef = await _get_recur_def(company)
    if period not in [p["key"] for p in rdef["periods"]]:
        raise HTTPException(status_code=400, detail="Periodo non valido")

    members = await db.recurrence_members.find({"company": company}, {"_id": 0}).to_list(20000)
    client_ids = [m["client_id"] for m in members]
    clients = await db.clients.find(
        {"id": {"$in": client_ids}, "deleted_at": None}, {"_id": 0}
    ).to_list(20000)
    cmap = {c["id"]: c for c in clients}

    is_admin = user.get("role") == "admin"
    visible = []
    for m in members:
        cli = cmap.get(m["client_id"])
        if not cli:
            continue  # client deleted
        if not is_admin and cli.get("agent") != user["username"]:
            continue
        visible.append((m, cli))

    done = await _recurrence_done_ids(company, period, [c["id"] for _, c in visible])

    # Group by recurrence group, preserving def group order then position.
    group_order = {g: i for i, g in enumerate(rdef["groups"])}
    visible.sort(key=lambda mc: (group_order.get(mc[0]["group"], 999),
                                 mc[0].get("position", 999),
                                 mc[1].get("ragione_sociale", "").lower()))
    groups_out = []
    current = None
    for m, cli in visible:
        if current is None or current["group"] != m["group"]:
            current = {"group": m["group"], "clients": []}
            groups_out.append(current)
        pub = client_public(cli)
        pub["recurrence_status"] = "ordine_effettuato" if cli["id"] in done else "da_gestire"
        pub["order_date"] = done.get(cli["id"])
        pub["member_id"] = m["id"]
        current["clients"].append(pub)

    return {"company": company, "period": period, "groups": groups_out}


@api.post("/recurrences/{company}/order")
async def recurrence_order(company: str, body: RecurrenceOrderBody, user=Depends(get_current_user)):
    rdef = await _get_recur_def(company)
    if body.period not in [p["key"] for p in rdef["periods"]]:
        raise HTTPException(status_code=400, detail="Periodo non valido")
    cli = await db.clients.find_one({"id": body.client_id, "deleted_at": None}, {"_id": 0})
    if not cli:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    if user.get("role") != "admin" and cli.get("agent") != user["username"]:
        raise HTTPException(status_code=403, detail="Non puoi gestire questo cliente")
    now = now_utc()
    doc = {
        "id": str(uuid.uuid4()),
        "client_id": body.client_id,
        "type": "recurrence_order",
        "company_id": None,
        "company_name": None,
        "note_text": None,
        "reschedule_until": None,
        "recurrence_company": company,
        "recurrence_period": body.period,
        "agent": user["username"],
        "created_at": now,
        "deleted_at": None,
    }
    await db.events.insert_one(doc)
    return {"ok": True, "order_date": iso(now)}


@api.post("/recurrences/{company}/order/undo")
async def recurrence_order_undo(company: str, body: RecurrenceOrderBody, user=Depends(get_current_user)):
    cli = await db.clients.find_one({"id": body.client_id, "deleted_at": None}, {"_id": 0})
    if not cli:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    if user.get("role") != "admin" and cli.get("agent") != user["username"]:
        raise HTTPException(status_code=403, detail="Non puoi gestire questo cliente")
    start = _rome_year_start()
    await db.events.update_many(
        {"type": "recurrence_order", "recurrence_company": company,
         "recurrence_period": body.period, "client_id": body.client_id,
         "created_at": {"$gte": start}, "deleted_at": None},
        {"$set": {"deleted_at": now_utc()}},
    )
    return {"ok": True}


@api.post("/recurrences/{company}/members")
async def add_recurrence_member(company: str, body: RecurrenceMemberCreate, user=Depends(get_current_user)):
    rdef = await _get_recur_def(company)
    if body.group not in rdef["groups"]:
        raise HTTPException(status_code=400, detail="Gruppo non valido")

    agent = user["username"]
    if user.get("role") == "admin" and body.agent in VALID_AGENTS:
        agent = body.agent

    if body.client_id:
        cli = await db.clients.find_one({"id": body.client_id, "deleted_at": None}, {"_id": 0})
        if not cli:
            raise HTTPException(status_code=404, detail="Cliente non trovato")
        if user.get("role") != "admin" and cli.get("agent") != user["username"]:
            raise HTTPException(status_code=403, detail="Non puoi aggiungere questo cliente")
        client_id = cli["id"]
        agent = cli.get("agent") or agent
    else:
        if not (body.ragione_sociale and body.ragione_sociale.strip()):
            raise HTTPException(status_code=400, detail="Ragione sociale obbligatoria")
        client_id = str(uuid.uuid4())
        await db.clients.insert_one({
            "id": client_id,
            "ragione_sociale": body.ragione_sociale.strip(),
            "codice_azienda": "",
            "provincia": normalize_provincia(body.provincia),
            "giro_id": None,
            "position": 999,
            "citta": (body.citta or "").strip(),
            "zona": (body.citta or "").strip(),
            "indirizzo": (body.indirizzo or "").strip(),
            "cap": (body.cap or "").strip(),
            "telefono": (body.telefono or "").strip(),
            "email": (body.email or "").strip(),
            "agent": agent,
            "permanent_note": "",
            "last_visit_at": None,
            "snoozed_until": None,
            "extra": {"recurrence_only": True, "needs_review": False},
            "created_by": user["username"],
            "deleted_at": None,
            "created_at": now_utc(),
        })

    dup = await db.recurrence_members.find_one({"company": company, "client_id": client_id})
    if dup:
        raise HTTPException(status_code=400, detail="Cliente già presente in questa ricorrenza")

    last = await db.recurrence_members.find(
        {"company": company, "group": body.group}, {"_id": 0, "position": 1}
    ).sort("position", -1).to_list(1)
    pos = (last[0]["position"] + 1) if last else 0
    await db.recurrence_members.insert_one({
        "id": str(uuid.uuid4()),
        "company": company,
        "client_id": client_id,
        "group": body.group,
        "position": pos,
        "agent": agent,
        "created_at": now_utc(),
    })
    return {"ok": True, "client_id": client_id}


@api.delete("/recurrences/{company}/members/{member_id}")
async def remove_recurrence_member(company: str, member_id: str, user=Depends(get_current_user)):
    member = await db.recurrence_members.find_one({"id": member_id, "company": company}, {"_id": 0})
    if not member:
        raise HTTPException(status_code=404, detail="Voce non trovata")
    if user.get("role") != "admin":
        cli = await db.clients.find_one({"id": member["client_id"]}, {"_id": 0})
        if not cli or cli.get("agent") != user["username"]:
            raise HTTPException(status_code=403, detail="Non puoi rimuovere questo cliente")
    await db.recurrence_members.delete_one({"id": member_id, "company": company})
    return {"ok": True}


# ---------------------------------------------------------------------------
# Monthly statistics + Excel export (per giro)
# ---------------------------------------------------------------------------
MONTHS_IT = ["Gen", "Feb", "Mar", "Apr", "Mag", "Giu", "Lug", "Ago", "Set", "Ott", "Nov", "Dic"]


async def _monthly_grid(giro_id: str, year: int, user: dict):
    giro = await db.giri.find_one({"id": giro_id}, {"_id": 0})
    if not giro:
        raise HTTPException(status_code=404, detail="Giro non trovato")

    q = {"giro_id": giro_id, "deleted_at": None}
    if user.get("role") != "admin":
        q["agent"] = user["username"]
    clients = await db.clients.find(q, {"_id": 0}).to_list(5000)
    clients.sort(key=lambda d: (d.get("position", 999), d.get("ragione_sociale", "").lower()))
    client_ids = [c["id"] for c in clients]

    year_start = datetime(year, 1, 1, tzinfo=timezone.utc)
    year_end = datetime(year + 1, 1, 1, tzinfo=timezone.utc)
    events = await db.events.find(
        {"client_id": {"$in": client_ids}, "deleted_at": None,
         "created_at": {"$gte": year_start, "$lt": year_end}},
        {"_id": 0},
    ).to_list(100000)

    grid = {cid: {m: {"visit": False, "orders": [], "collection": False} for m in range(12)} for cid in client_ids}
    for e in events:
        created = e["created_at"]
        if created.tzinfo is None:
            created = created.replace(tzinfo=timezone.utc)
        m = created.astimezone(ROME).month - 1
        cell = grid[e["client_id"]][m]
        if e["type"] == "visit":
            cell["visit"] = True
        elif e["type"] == "order":
            name = e.get("company_name") or "Ordine"
            if name not in cell["orders"]:
                cell["orders"].append(name)
        elif e["type"] == "collection":
            cell["collection"] = True
    return giro, clients, grid


@api.get("/stats/monthly")
async def stats_monthly(giro_id: str = Query(...), year: int = Query(...), user=Depends(get_current_user)):
    giro, clients, grid = await _monthly_grid(giro_id, year, user)
    rows = []
    for cli in clients:
        months = []
        for m in range(12):
            c = grid[cli["id"]][m]
            months.append({"visit": c["visit"], "orders": c["orders"], "collection": c["collection"]})
        rows.append({
            "id": cli["id"],
            "ragione_sociale": cli.get("ragione_sociale", ""),
            "citta": cli.get("citta", ""),
            "months": months,
        })
    return {"giro": {"id": giro["id"], "name": giro["name"]}, "year": year, "months": MONTHS_IT, "clients": rows}


@api.get("/export/monthly")
async def export_monthly(giro_id: str = Query(...), year: int = Query(...), user=Depends(get_current_user)):
    giro, clients, grid = await _monthly_grid(giro_id, year, user)

    def mark_text(cell: dict) -> str:
        out = []
        if cell["visit"]:
            out.append("\u2713 visita")
        if cell["orders"]:
            out.append("\u2713 ordine: " + ", ".join(cell["orders"]))
        if cell["collection"]:
            out.append("\u2713 incasso")
        return "\n".join(out)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Riepilogo"

    header_fill = PatternFill("solid", fgColor="047857")
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14)

    ws.cell(row=1, column=1, value=f"Giro: {giro['name']} \u2014 Anno {year}").font = title_font
    ws.append([])

    headers = ["Cliente", "Citt\u00e0"] + MONTHS_IT
    ws.append(headers)
    hrow = ws.max_row
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=hrow, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    for cli in clients:
        row = [cli.get("ragione_sociale", ""), cli.get("citta", "")]
        for m in range(12):
            row.append(mark_text(grid[cli["id"]][m]))
        ws.append(row)
        r = ws.max_row
        for col in range(3, 15):
            ws.cell(row=r, column=col).alignment = Alignment(horizontal="center", wrap_text=True, vertical="top")

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 20
    for i in range(12):
        ws.column_dimensions[openpyxl.utils.get_column_letter(3 + i)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    safe = "".join(ch for ch in giro["name"] if ch.isalnum() or ch in " -_").strip().replace(" ", "_")
    filename = f"Riepilogo_{safe}_{year}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------------------------------------------------------------------------
# Integrazione GooseApp (sincronizzazione ordini)
# ---------------------------------------------------------------------------
import gooseapp  # noqa: E402

# Mappatura campi configurabile: i nomi tabella/campi reali di GooseApp si
# scoprono con TABLES/STRUCTURE (dopo whitelist IP su Render) e si salvano qui.
GOOSE_DEFAULT_CONFIG = {
    "orders_table": "",       # es. "DOC_TES" (da scoprire)
    "field_order_id": "RECORD_ID",
    "field_piva": "",         # campo con la Partita IVA del cliente
    "field_site": "",         # campo con il codice sede/destinazione
    "field_company": "",      # campo con il mandante/azienda
    "field_number": "",       # numero documento
    "field_date": "",         # data ordine (yyyy-mm-dd ...)
    "field_ship": "",         # campo/i indirizzo di spedizione o destinazione (separati da virgola)
    "dest_table": "",         # tabella destinazioni (se field_ship contiene un ID)
    "mandanti_table": "",     # tabella mandanti GooseApp (RECORD_ID = IDMAN)
    "field_payment": "",      # campo modalità di pagamento dell'ordine
    "payment_table": "",      # tabella modalità di pagamento (se il campo contiene un ID)
}


class GooseConfigBody(BaseModel):
    orders_table: Optional[str] = None
    field_order_id: Optional[str] = None
    field_piva: Optional[str] = None
    field_site: Optional[str] = None
    field_company: Optional[str] = None
    field_number: Optional[str] = None
    field_date: Optional[str] = None
    field_ship: Optional[str] = None
    dest_table: Optional[str] = None
    mandanti_table: Optional[str] = None
    field_payment: Optional[str] = None
    payment_table: Optional[str] = None


class GooseMandanteBody(BaseModel):
    code: str
    company_id: Optional[str] = None


class GooseSyncBody(BaseModel):
    # from GooseApp incrementale: "latest" (default) o date/time "yyyy-mm-dd hh:mm:ss"
    from_: Optional[str] = Field(default="latest", alias="from")
    # SOLO admin/anteprima: record simulati per testare il mapping senza API live
    mock_records: Optional[List[dict]] = None

    class Config:
        populate_by_name = True


async def _goose_config() -> dict:
    doc = await db.goose_config.find_one({"id": "config"}, {"_id": 0}) or {}
    cfg = dict(GOOSE_DEFAULT_CONFIG)
    cfg.update({k: v for k, v in doc.items() if k in GOOSE_DEFAULT_CONFIG and v})
    return cfg


def _norm_piva(raw) -> str:
    """Normalizza una Partita IVA per il confronto: solo alfanumerici, maiuscolo,
    rimuove un eventuale prefisso paese 'IT'."""
    s = re.sub(r"[^A-Za-z0-9]", "", str(raw or "")).upper()
    if s.startswith("IT") and len(s) > 11:
        s = s[2:]
    return s


def _parse_goose_date(raw) -> Optional[datetime]:
    if not raw:
        return None
    s = str(raw).strip().replace("T", " ")
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(s[:len(fmt) + 2], fmt).replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    return None


def _clean_doc_number(raw) -> str:
    """Ricava il numero ordine/documento 'pulito' (es. DTNUM) nascondendo il
    RECORD_ID interno. Se il valore è composito (es. '2504111729051590-143' o
    '...–143'), prende l'ultimo segmento dopo il trattino."""
    s = str(raw or "").strip()
    if not s:
        return ""
    s = s.replace("\u2013", "-").replace("\u2014", "-")  # en/em dash -> '-'
    if "-" in s:
        s = s.rsplit("-", 1)[-1].strip()
    return s


def _match_site(cli: dict, site_ref: str) -> dict:
    """Associa il codice sede dell'ordine GooseApp a una sede del cliente.
    Fallback: sede principale."""
    extra = cli.get("extra") or {}
    ref = (site_ref or "").strip()
    if ref:
        for s in extra.get("sites", []):
            if (s.get("goose_code") or "").strip() and s.get("goose_code").strip() == ref:
                return {"site_id": s.get("id"), "site_nominativo": s.get("nominativo") or "",
                        "site_indirizzo": s.get("indirizzo") or "", "matched": True, "by_code": True}
        # match sul codice della sede principale
        if (extra.get("goose_code") or "").strip() == ref:
            return {"site_id": "principal", "site_nominativo": extra.get("sede_nome") or "",
                    "site_indirizzo": cli.get("indirizzo") or "", "matched": True, "by_code": True}
    # default: sede principale
    return {"site_id": "principal", "site_nominativo": extra.get("sede_nome") or "",
            "site_indirizzo": cli.get("indirizzo") or "", "matched": not ref}


_NAME_KEYS = ("RAGSO", "RAGIONE", "DENOM", "NOME", "DESCR", "DES", "NOM", "RAG")
_CITY_KEYS = ("LOCALITA", "CITTA", "COMUNE", "LOC", "CIT", "COM")


def _record_text(rec: dict, prefs) -> str:
    """Primo valore testuale (non numerico/ID) di un record GooseApp il cui nome
    campo contiene una delle chiavi preferite (in ordine di priorità)."""
    items = [(str(k).upper(), v) for k, v in (rec or {}).items()]
    for pref in prefs:
        for k, v in items:
            if k.startswith("ID") or k == "RECORD_ID" or pref not in k:
                continue
            t = str(v or "").strip()
            if t and not re.fullmatch(r"[\d\-./ ]+", t):
                return t
    return ""


def _match_company(name: str, companies: list):
    """Azienda in anagrafica il cui nome compare (a parole intere) nel testo."""
    n = f" {_norm(name)} "
    best = None
    for c in companies:
        cn = _norm(c.get("name"))
        if cn and (f" {cn} " in n or n.strip() == cn):
            if not best or len(cn) > len(_norm(best["name"])):
                best = c
    return best


def _ship_label(text: str) -> str:
    """Dal testo dell'indirizzo di spedizione ricava una dicitura breve: la
    città (dopo il CAP) se presente, altrimenti la prima riga/denominazione."""
    t = re.sub(r"\s+", " ", str(text or "")).strip()
    if not t:
        return ""
    m = re.search(r"\b\d{5}\s+([A-Za-zÀ-ÿ' .]+?)\s*(?:\(|,|-|$)", t)
    if m and m.group(1).strip():
        return m.group(1).strip().title()
    t = re.sub(r"^(sede( di scarico)?|destinazione|punto vendita)\s*[:\-]?\s*", "", t.split(",")[0].strip(), flags=re.I)
    return t[:60]


async def _goose_lookup(table: str) -> dict:
    """RECORD_ID -> record di una tabella GooseApp (per IDMAN / destinazioni)."""
    if not table:
        return {}
    try:
        rows = await gooseapp.select(table)
    except gooseapp.GooseError as e:
        logger.warning("Lookup GooseApp %s non riuscito: %s", table, e)
        return {}
    return {str(_ci_get(r, "RECORD_ID") or "").strip(): r for r in rows if _ci_get(r, "RECORD_ID")}


async def _process_goose_orders(records: list, cfg: dict, actor: dict) -> dict:
    """Associa ciascun ordine al cliente (per P.IVA) e alla sede, creando un
    evento 'order' con source='gooseapp'. Salta i duplicati (goose_order_id)."""
    if not cfg.get("field_piva"):
        raise HTTPException(status_code=400, detail="Configurazione incompleta: manca il campo Partita IVA (field_piva).")

    # Indice P.IVA -> cliente (solo clienti con P.IVA valorizzata)
    piva_map = {}
    async for c in db.clients.find({"deleted_at": None}, {"_id": 0}):
        piva = _norm_piva((c.get("extra") or {}).get("partita_iva"))
        if piva:
            piva_map.setdefault(piva, c)

    # Mandanti: il valore GooseApp (IDMAN/azienda) viene ricondotto al nome
    # dell'azienda in anagrafica quando coincide (match normalizzato).
    companies = await db.companies.find({}, {"_id": 0}).to_list(1000)
    comp_by_id = {c["id"]: c for c in companies}
    comp_map = {}
    for c in companies:
        comp_map[_norm(c["name"])] = c
        if c.get("goose_code"):
            comp_map[_norm(c["goose_code"])] = c
    # Tabella di conversione permanente IDMAN -> azienda
    man_map = {m["code"]: m async for m in db.goose_mandanti.find({}, {"_id": 0})}
    man_lookup = await _goose_lookup(cfg.get("mandanti_table")) if cfg.get("mandanti_table") else {}
    dest_lookup = await _goose_lookup(cfg.get("dest_table")) if cfg.get("dest_table") else {}
    ship_fields = [f.strip() for f in (cfg.get("field_ship") or "").split(",") if f.strip()]
    pay_lookup = await _goose_lookup(cfg.get("payment_table")) if cfg.get("payment_table") else {}

    is_admin = actor.get("role") == "admin"
    imported, skipped, unmatched, errors = 0, 0, [], []
    details = []

    for rec in records:
        oid = str(rec.get(cfg["field_order_id"]) or "").strip()
        piva = _norm_piva(rec.get(cfg["field_piva"]))
        try:
            if oid:
                dup = await db.events.find_one({"goose_order_id": oid, "deleted_at": None})
                if dup:
                    skipped += 1
                    continue
            cli = piva_map.get(piva)
            if not cli:
                unmatched.append({"order_id": oid, "partita_iva": piva, "reason": "P.IVA non trovata in anagrafica"})
                continue
            # un agente importa solo gli ordini dei propri clienti
            if not is_admin and cli.get("agent") != actor["username"]:
                continue
            site = _match_site(cli, str(rec.get(cfg.get("field_site") or "") or ""))
            created = _parse_goose_date(rec.get(cfg.get("field_date") or "")) or now_utc()
            company_code = str(rec.get(cfg.get("field_company") or "") or "").strip()
            company_id, company_name = await _resolve_mandante(
                company_code, man_map, man_lookup, comp_map, comp_by_id, companies)
            ship_label = ""
            if ship_fields:
                raw_ship = " ".join(str(rec.get(f) or "").strip() for f in ship_fields).strip()
                dest = dest_lookup.get(raw_ship)
                if dest:
                    ship_label = (_record_text(dest, _CITY_KEYS) or _record_text(dest, _NAME_KEYS)).title()
                elif raw_ship and not re.fullmatch(r"[\d\-]+", raw_ship):
                    ship_label = _ship_label(raw_ship)
                if ship_label and not site.get("by_code"):
                    # se la sede del cliente ha la stessa città/nome, la associamo
                    k = _norm(ship_label)
                    for st in (cli.get("extra") or {}).get("sites", []):
                        if k and k in (_norm(st.get("citta")), _norm(st.get("nominativo"))):
                            site = {"site_id": st.get("id"), "site_nominativo": st.get("nominativo") or "",
                                    "site_indirizzo": st.get("indirizzo") or "", "matched": True}
                            break
            order_number = _clean_doc_number(rec.get(cfg.get("field_number"))) if cfg.get("field_number") else ""
            pay_raw = str(rec.get(cfg.get("field_payment") or "") or "").strip()
            pay_rec = pay_lookup.get(pay_raw)
            pay_text = (_record_text(pay_rec, _NAME_KEYS) or " ".join(str(v) for v in pay_rec.values())) if pay_rec else pay_raw
            pay_mode = payment_mode_from_text(pay_text)
            pay_days = PAYMENT_MODES[pay_mode]["days"] if pay_mode else None
            doc = {
                "id": str(uuid.uuid4()),
                "client_id": cli["id"],
                "type": "order",
                "company_id": company_id,
                "company_name": company_name,
                "note_text": None,
                "order_number": order_number or None,
                "reschedule_until": None,
                "payment_mode": pay_mode,
                "payment_text": pay_text or None,
                "due_at": (created + timedelta(days=pay_days)) if pay_days else None,
                "source": "gooseapp",
                "goose_order_id": oid or None,
                "goose_company_code": company_code or None,
                "ship_label": ship_label or None,
                "goose_record": {str(k): (v if isinstance(v, (str, int, float, bool)) or v is None else str(v))
                                 for k, v in rec.items()},
                "collection_method": None,
                "collection_ref_date": None,
                "site_id": site["site_id"],
                "site_nominativo": site["site_nominativo"],
                "site_indirizzo": site["site_indirizzo"],
                "agent": cli.get("agent"),
                "created_at": created,
                "deleted_at": None,
            }
            await db.events.insert_one(doc)
            imported += 1
            details.append({"order_id": oid, "client": cli.get("ragione_sociale"),
                            "site": site["site_nominativo"] or "Sede principale",
                            "site_matched": site["matched"]})
        except Exception as e:  # noqa: BLE001
            logger.error("Goose order %s errore: %s", oid, e)
            errors.append({"order_id": oid, "error": str(e)})

    summary = {
        "total": len(records),
        "imported": imported,
        "skipped_duplicates": skipped,
        "unmatched": unmatched,
        "errors": errors,
        "details": details,
        "synced_at": iso(now_utc()),
    }
    await db.goose_config.update_one(
        {"id": "config"},
        {"$set": {"id": "config", "last_sync_at": now_utc(), "last_result": summary}},
        upsert=True,
    )
    return summary


async def _resolve_mandante(code, man_map, man_lookup, comp_map, comp_by_id, companies):
    """IDMAN -> (company_id, company_name). Ordine: tabella di conversione
    salvata, nome/codice azienda, tabella mandanti GooseApp (auto-salvata)."""
    if not code:
        return None, None
    m = man_map.get(code)
    if m and m.get("company_id") and comp_by_id.get(m["company_id"]):
        c = comp_by_id[m["company_id"]]
        return c["id"], c["name"]
    c = comp_map.get(_norm(code))
    if c:
        return c["id"], c["name"]
    if m and m.get("goose_name"):
        c = _match_company(m["goose_name"], companies)
        return (c["id"], c["name"]) if c else (None, m["goose_name"])
    rec = man_lookup.get(code)
    if rec:
        gname = _record_text(rec, _NAME_KEYS)
        c = _match_company(gname, companies) if gname else None
        if not c:
            for v in rec.values():
                c = _match_company(str(v or ""), companies)
                if c:
                    break
        entry = {"code": code, "goose_name": gname or None,
                 "company_id": c["id"] if c else None, "updated_at": now_utc(), "auto": True}
        await db.goose_mandanti.update_one({"code": code}, {"$set": entry}, upsert=True)
        man_map[code] = entry
        if c:
            return c["id"], c["name"]
        if gname:
            return None, gname
    return None, code


async def _apply_mandante_to_events(code: str, company: Optional[dict], goose_name: Optional[str] = None):
    q = {"source": "gooseapp", "$or": [{"goose_company_code": code}, {"company_name": code, "company_id": None}]}
    if company:
        st = {"company_id": company["id"], "company_name": company["name"], "goose_company_code": code}
    else:
        st = {"company_id": None, "company_name": goose_name or code, "goose_company_code": code}
    res = await db.events.update_many(q, {"$set": st})
    return res.modified_count


@api.get("/goose/mandanti")
async def goose_mandanti(user=Depends(require_admin)):
    """Tabella di conversione IDMAN -> azienda, con i codici visti negli ordini."""
    companies = {c["id"]: c for c in await db.companies.find({}, {"_id": 0}).to_list(1000)}
    maps = {m["code"]: m async for m in db.goose_mandanti.find({}, {"_id": 0})}
    counts, ev_comp = {}, {}
    async for e in db.events.find({"source": "gooseapp", "deleted_at": None},
                                  {"_id": 0, "goose_company_code": 1, "company_name": 1, "company_id": 1}):
        code = e.get("goose_company_code") or (e.get("company_name") if not e.get("company_id") else None)
        if code:
            counts[code] = counts.get(code, 0) + 1
            if e.get("company_id"):
                ev_comp.setdefault(code, e["company_id"])
    out = []
    for code in sorted(set(counts) | set(maps)):
        m = maps.get(code, {})
        c = companies.get(m.get("company_id") or ev_comp.get(code))
        out.append({"code": code, "orders": counts.get(code, 0), "goose_name": m.get("goose_name"),
                    "company_id": c["id"] if c else None, "company_name": c["name"] if c else None})
    out.sort(key=lambda r: (r["company_id"] is not None, -r["orders"]))
    return out


@api.put("/goose/mandanti")
async def goose_set_mandante(body: GooseMandanteBody, user=Depends(require_admin)):
    code = body.code.strip()
    if not code:
        raise HTTPException(status_code=400, detail="Codice mandante obbligatorio")
    company = None
    if body.company_id:
        company = await db.companies.find_one({"id": body.company_id}, {"_id": 0})
        if not company:
            raise HTTPException(status_code=404, detail="Azienda non trovata")
    prev = await db.goose_mandanti.find_one({"code": code}, {"_id": 0}) or {}
    await db.goose_mandanti.update_one(
        {"code": code},
        {"$set": {"code": code, "company_id": company["id"] if company else None,
                  "updated_at": now_utc(), "auto": False}},
        upsert=True)
    n = await _apply_mandante_to_events(code, company, prev.get("goose_name"))
    return {"ok": True, "updated_orders": n}


@api.post("/goose/mandanti/auto")
async def goose_auto_mandanti(user=Depends(require_admin)):
    """Rileva automaticamente i nomi dei mandanti dalla tabella mandanti di
    GooseApp (configurata o cercata tra le tabelle con 'MAN' nel nome) e
    aggiorna la tabella di conversione + gli ordini già importati."""
    cfg = await _goose_config()
    codes = {r["code"] for r in await goose_mandanti(user) if not r["company_id"]}
    candidates = [cfg["mandanti_table"]] if cfg.get("mandanti_table") else []
    if not candidates:
        try:
            tabs = _normalize_tables(await gooseapp.list_tables())
        except gooseapp.GooseError as e:
            raise HTTPException(status_code=502, detail=str(e))
        candidates = [t["name"] for t in tabs if "MAN" in t["name"].upper()]
    companies = await db.companies.find({}, {"_id": 0}).to_list(1000)
    used, mapped, named, unresolved = None, [], [], []
    for table in candidates:
        lookup = await _goose_lookup(table)
        hits = {c for c in codes if c in lookup}
        if not hits:
            continue
        used = table
        for code in hits:
            rec = lookup[code]
            gname = _record_text(rec, _NAME_KEYS)
            c = _match_company(gname, companies) if gname else None
            if not c:
                for v in rec.values():
                    c = _match_company(str(v or ""), companies)
                    if c:
                        break
            await db.goose_mandanti.update_one(
                {"code": code},
                {"$set": {"code": code, "goose_name": gname or None, "company_id": c["id"] if c else None,
                          "updated_at": now_utc(), "auto": True}}, upsert=True)
            await _apply_mandante_to_events(code, c, gname)
            (mapped if c else named).append({"code": code, "goose_name": gname, "company": c["name"] if c else None})
        break
    if used and not cfg.get("mandanti_table"):
        await db.goose_config.update_one({"id": "config"}, {"$set": {"mandanti_table": used}}, upsert=True)
    unresolved = sorted(codes - {m["code"] for m in mapped + named})
    return {"table": used, "candidates": candidates, "mapped": mapped, "named_only": named, "unresolved": unresolved}


@api.get("/goose/status")
async def goose_status(user=Depends(get_current_user)):
    doc = await db.goose_config.find_one({"id": "config"}, {"_id": 0}) or {}
    cfg = await _goose_config()
    return {
        "configured": gooseapp.is_configured(),
        "ready": bool(cfg.get("orders_table") and cfg.get("field_piva")),
        "last_sync_at": iso(doc.get("last_sync_at")),
        "last_result": doc.get("last_result"),
    }


@api.get("/goose/config")
async def goose_get_config(user=Depends(require_admin)):
    return await _goose_config()


@api.put("/goose/config")
async def goose_put_config(body: GooseConfigBody, user=Depends(require_admin)):
    update = {k: v for k, v in body.dict(exclude_unset=True).items() if v is not None}
    update["id"] = "config"
    await db.goose_config.update_one({"id": "config"}, {"$set": update}, upsert=True)
    return await _goose_config()


def _ci_get(d: dict, *keys):
    """Lettura case-insensitive: ritorna il primo valore non vuoto tra le chiavi date."""
    lower = {str(k).lower(): v for k, v in d.items()}
    for k in keys:
        v = lower.get(str(k).lower())
        if v not in (None, ""):
            return v
    return None


def _normalize_tables(result) -> list:
    """Estrae una lista uniforme [{name, description}] dalla risposta TABLES,
    tollerante a formati diversi (lista di stringhe, lista di dict, dict)."""
    data = result.get("data") if isinstance(result, dict) else result
    if not isinstance(data, list):
        if isinstance(result, dict):
            data = [{"name": k, "description": v} for k, v in result.items()
                    if k not in ("total", "data", "status")]
        else:
            data = []
    out = []
    for item in data:
        if isinstance(item, str):
            out.append({"name": item, "description": ""})
        elif isinstance(item, dict):
            name = _ci_get(item, "table", "name", "code", "id", "tabella", "field") or ""
            desc = _ci_get(item, "description", "label", "desc", "descrizione", "nome") or ""
            if name:
                out.append({"name": str(name), "description": str(desc)})
    return out


def _normalize_fields(result) -> list:
    """Estrae una lista uniforme [{name, type}] dalla risposta STRUCTURE.
    GooseApp (DOC_TES) usa result.data con oggetti {"Field": "...", ...}."""
    data = None
    if isinstance(result, dict):
        data = (result.get("data") or result.get("fields") or result.get("structure")
                or result.get("campi"))
    if data is None:
        data = result if isinstance(result, list) else []
    out = []
    if isinstance(data, dict):
        data = [{"name": k, "type": v} for k, v in data.items()]
    for item in (data or []):
        if isinstance(item, str):
            out.append({"name": item, "type": ""})
        elif isinstance(item, dict):
            name = _ci_get(item, "Field", "field", "name", "code", "id", "campo") or ""
            typ = _ci_get(item, "type", "tipo", "datatype", "Type") or ""
            if name:
                out.append({"name": str(name), "type": str(typ)})
    return out


@api.get("/goose/tables")
async def goose_tables(structure: bool = False, user=Depends(require_admin)):
    """Scoperta schema: elenco tabelle GooseApp (richiede IP in whitelist)."""
    try:
        result = await gooseapp.list_tables(with_structure=structure)
        return {"tables": _normalize_tables(result), "raw": result}
    except gooseapp.GooseError as e:
        raise HTTPException(status_code=502, detail=str(e))


@api.get("/goose/structure/{table}")
async def goose_structure(table: str, user=Depends(require_admin)):
    try:
        result = await gooseapp.get_structure(table)
        return {"fields": _normalize_fields(result), "raw": result}
    except gooseapp.GooseError as e:
        raise HTTPException(status_code=502, detail=str(e))


@api.post("/goose/sync")
async def goose_sync(body: GooseSyncBody = GooseSyncBody(), user=Depends(get_current_user)):
    """Preleva gli ordini da GooseApp e li importa associandoli per P.IVA + sede.
    Se `mock_records` è fornito (solo admin) salta la chiamata live: utile per
    testare il mapping in anteprima senza IP in whitelist."""
    cfg = await _goose_config()
    if body.mock_records is not None:
        if user.get("role") != "admin":
            raise HTTPException(status_code=403, detail="Solo admin può usare mock_records")
        records = body.mock_records
    else:
        if not cfg.get("orders_table"):
            raise HTTPException(status_code=400, detail="Configura prima la tabella Ordini (orders_table).")
        try:
            records = await gooseapp.select(cfg["orders_table"], from_=body.from_ or "latest")
        except gooseapp.GooseError as e:
            raise HTTPException(status_code=502, detail=str(e))
    return await _process_goose_orders(records, cfg, user)


app.include_router(api)


# CORS: elenco esplicito delle origini consentite (il wildcard "*" con
# allow_credentials=True viene rifiutato dai browser). Copriamo il dominio
# Vercel di produzione, i deploy di anteprima *.vercel.app, l'anteprima
# Emergent e lo sviluppo locale.
ALLOWED_ORIGINS = [
    "https://rodomisto-rappresentanze-frontend.vercel.app",
    "http://localhost:3000",
    "http://localhost:8081",
    "http://localhost:19006",
]
app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_origin_regex=r"https://([a-z0-9-]+\.)*(vercel\.app|emergentagent\.com)$",
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
