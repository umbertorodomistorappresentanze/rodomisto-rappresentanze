import io
import logging
import os
import re
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import List, Optional
from zoneinfo import ZoneInfo

import bcrypt
import jwt
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from dotenv import load_dotenv
from fastapi import APIRouter, Depends, FastAPI, HTTPException, Query
from fastapi.responses import StreamingResponse
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field
from starlette.middleware.cors import CORSMiddleware

import seed_data
import recurrence_seed
import unicodedata

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(name)s - %(levelname)s - %(message)s")
logger = logging.getLogger("agendavisite")

mongo_url = os.environ["MONGO_URL"]
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ["DB_NAME"]]

JWT_SECRET = os.environ["JWT_SECRET"]
JWT_ALGO = "HS256"
ACCESS_TOKEN_DAYS = int(os.environ.get("ACCESS_TOKEN_DAYS", "30"))
ROME = ZoneInfo("Europe/Rome")
VISIT_THRESHOLD_DAYS = 21

# Payment modes selectable per single order. days=None means no future suspension.
PAYMENT_MODES = {
    "anticipato": {"label": "Anticipato", "days": None},
    "contrassegno": {"label": "Contrassegno", "days": None},
    "bonifico_30": {"label": "Bonifico bancario 30 giorni", "days": 30},
    "bonifico_60": {"label": "Bonifico bancario 60 giorni", "days": 60},
    "agente_30": {"label": "Pagamento mezzo Agente 30 giorni", "days": 30},
    "agente_60": {"label": "Pagamento mezzo Agente 60 giorni", "days": 60},
    "agente_90": {"label": "Pagamento mezzo Agente 90 giorni", "days": 90},
    "rifatturazione_pac": {"label": "Rifatturazione Pac", "days": None},
}
PAYMENT_MODE_ORDER = ["anticipato", "contrassegno", "bonifico_30", "bonifico_60", "agente_30", "agente_60", "agente_90", "rifatturazione_pac"]

app = FastAPI(title="Rodomisto Rappresentanze API")
api = APIRouter(prefix="/api")
bearer = HTTPBearer(auto_error=False)


@app.get("/health")
async def health():
    return {"status": "ok"}


@api.get("/health")
async def api_health():
    return {"status": "ok"}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def now_utc() -> datetime:
    return datetime.now(timezone.utc)


def start_of_today_utc() -> datetime:
    local = datetime.now(ROME)
    start_local = local.replace(hour=0, minute=0, second=0, microsecond=0)
    return start_local.astimezone(timezone.utc)


def hash_pw(pw: str) -> str:
    return bcrypt.hashpw(pw.encode("utf-8")[:72], bcrypt.gensalt()).decode("utf-8")


def verify_pw(pw: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(pw.encode("utf-8")[:72], hashed.encode("utf-8"))
    except Exception:
        return False


def create_token(username: str) -> str:
    now = now_utc()
    payload = {"sub": username, "iat": now, "exp": now + timedelta(days=ACCESS_TOKEN_DAYS)}
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGO)


def iso(dt: Optional[datetime]) -> Optional[str]:
    if dt is None:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc).isoformat()


async def get_current_user(creds: Optional[HTTPAuthorizationCredentials] = Depends(bearer)):
    unauth = HTTPException(status_code=401, detail="Credenziali non valide o scadute",
                           headers={"WWW-Authenticate": "Bearer"})
    if creds is None or creds.scheme.lower() != "bearer":
        raise unauth
    try:
        payload = jwt.decode(creds.credentials, JWT_SECRET, algorithms=[JWT_ALGO],
                             options={"require": ["sub", "exp"]})
        username = payload["sub"]
    except Exception:
        raise unauth
    user = await db.users.find_one({"username": username, "is_active": True}, {"_id": 0})
    if not user:
        raise unauth
    return user


# ---------------------------------------------------------------------------
# Models
# ---------------------------------------------------------------------------
class LoginRequest(BaseModel):
    username: str
    password: str


class Company(BaseModel):
    name: str


class CompanyUpdate(BaseModel):
    name: Optional[str] = None
    active: Optional[bool] = None


class GiroCreate(BaseModel):
    name: str
    localities: List[str] = Field(default_factory=list)


class GiroUpdate(BaseModel):
    name: Optional[str] = None
    localities: Optional[List[str]] = None
    active: Optional[bool] = None


class ClientCreate(BaseModel):
    ragione_sociale: str
    provincia: str = ""
    giro_id: Optional[str] = None
    position: Optional[int] = None
    citta: str = ""
    zona: str = ""
    indirizzo: str = ""
    cap: str = ""
    telefono: str = ""
    email: str = ""
    agent: Optional[str] = None


class ClientUpdate(BaseModel):
    ragione_sociale: Optional[str] = None
    provincia: Optional[str] = None
    giro_id: Optional[str] = None
    position: Optional[int] = None
    citta: Optional[str] = None
    zona: Optional[str] = None
    indirizzo: Optional[str] = None
    cap: Optional[str] = None
    telefono: Optional[str] = None
    email: Optional[str] = None
    permanent_note: Optional[str] = None
    agent: Optional[str] = None


class EventCreate(BaseModel):
    client_id: str
    type: str  # order | collection | reschedule | note | suspension | visit(legacy)
    company_id: Optional[str] = None
    note_text: Optional[str] = None
    reschedule_days: Optional[int] = None
    reschedule_date: Optional[str] = None  # ISO date string
    payment_mode: Optional[str] = None  # for orders


# ---------------------------------------------------------------------------
# Serializers
# ---------------------------------------------------------------------------
def client_public(doc: dict) -> dict:
    return {
        "id": doc["id"],
        "ragione_sociale": doc.get("ragione_sociale", ""),
        "provincia": doc.get("provincia", ""),
        "giro_id": doc.get("giro_id"),
        "position": doc.get("position"),
        "citta": doc.get("citta", ""),
        "zona": doc.get("zona", ""),
        "indirizzo": doc.get("indirizzo", ""),
        "cap": doc.get("cap", ""),
        "telefono": doc.get("telefono", ""),
        "email": doc.get("email", ""),
        "agent": doc.get("agent", ""),
        "permanent_note": doc.get("permanent_note", ""),
        "last_visit_at": iso(doc.get("last_visit_at")),
        "snoozed_until": iso(doc.get("snoozed_until")),
        "extra": doc.get("extra", {}),
    }


def event_public(doc: dict) -> dict:
    return {
        "id": doc["id"],
        "client_id": doc["client_id"],
        "type": doc["type"],
        "company_id": doc.get("company_id"),
        "company_name": doc.get("company_name"),
        "note_text": doc.get("note_text"),
        "reschedule_until": iso(doc.get("reschedule_until")),
        "agent": doc.get("agent"),
        "created_at": iso(doc.get("created_at")),
        "recurrence_company": doc.get("recurrence_company"),
        "recurrence_period": doc.get("recurrence_period"),
        "payment_mode": doc.get("payment_mode"),
        "payment_mode_label": PAYMENT_MODES.get(doc.get("payment_mode"), {}).get("label") if doc.get("payment_mode") else None,
        "due_at": iso(doc.get("due_at")),
    }


# ---------------------------------------------------------------------------
# Seeding / import
# ---------------------------------------------------------------------------
async def seed():
    await db.users.create_index("username", unique=True)
    seeds = [
        {"username": "umberto", "display_name": "Umberto Rodomisto", "role": "admin",
         "password": os.environ["SEED_UMBERTO_PASSWORD"]},
        {"username": "andrea", "display_name": "Andrea Azzarito", "role": "agent",
         "password": os.environ["SEED_ANDREA_PASSWORD"]},
    ]
    for s in seeds:
        existing = await db.users.find_one({"username": s["username"]})
        if not existing:
            await db.users.insert_one({
                "id": str(uuid.uuid4()),
                "username": s["username"],
                "display_name": s["display_name"],
                "role": s["role"],
                "hashed_password": hash_pw(s["password"]),
                "is_active": True,
                "created_at": now_utc(),
            })
        else:
            updates = {"role": s["role"]}
            if not verify_pw(s["password"], existing.get("hashed_password", "")):
                updates["hashed_password"] = hash_pw(s["password"])
            await db.users.update_one({"username": s["username"]}, {"$set": updates})

    for i, name in enumerate(seed_data.COMPANIES):
        existing = await db.companies.find_one({"name": name})
        if not existing:
            await db.companies.insert_one({
                "id": str(uuid.uuid4()), "name": name, "active": True,
                "order": i, "created_at": now_utc(),
            })

    giro_count = await db.giri.count_documents({})
    if giro_count == 0:
        for i, g in enumerate(seed_data.GIRI):
            await db.giri.insert_one({
                "id": str(uuid.uuid4()), "name": g["name"], "localities": g["localities"],
                "order": i, "active": True, "created_at": now_utc(),
            })

    client_count = await db.clients.count_documents({})
    if client_count == 0:
        giri_docs = await db.giri.find({}, {"_id": 0}).to_list(1000)
        loc_index = seed_data.build_locality_index(giri_docs)
        parsed = seed_data.parse_clients()
        docs = []
        deleted = 0
        for c in parsed:
            if seed_data.should_delete(c["ragione_sociale"], c["citta"]):
                deleted += 1
                continue
            giro_id, position = seed_data.resolve_assignment(
                c["ragione_sociale"], c["zona"], giri_docs, loc_index
            )
            docs.append({
                "id": str(uuid.uuid4()),
                "ragione_sociale": c["ragione_sociale"],
                "codice_azienda": c["codice_azienda"],
                "provincia": c["provincia"],
                "giro_id": giro_id,
                "position": position if position is not None else 999,
                "citta": c["citta"],
                "zona": c["zona"],
                "indirizzo": c["indirizzo"],
                "cap": c["cap"],
                "telefono": c["telefono"],
                "email": c["email"],
                "agent": c["agent"],
                "permanent_note": "",
                "last_visit_at": None,
                "snoozed_until": None,
                "extra": c["extra"],
                "created_by": c["agent"],
                "deleted_at": None,
                "created_at": now_utc(),
            })
        # Extra clients not present in the Excel.
        for ec in seed_data.EXTRA_CLIENTS:
            giro_id, position = seed_data.giro_position(giri_docs, ec["giro_name"], ec["locality"])
            docs.append({
                "id": str(uuid.uuid4()),
                "ragione_sociale": ec["ragione_sociale"],
                "codice_azienda": "",
                "provincia": ec.get("provincia", ""),
                "giro_id": giro_id,
                "position": position if position is not None else 999,
                "citta": ec.get("citta", ""),
                "zona": ec.get("zona", ""),
                "indirizzo": ec.get("indirizzo", ""),
                "cap": ec.get("cap", ""),
                "telefono": ec.get("telefono", ""),
                "email": ec.get("email", ""),
                "agent": ec.get("agent", "umberto"),
                "permanent_note": "",
                "last_visit_at": None,
                "snoozed_until": None,
                "extra": {},
                "created_by": ec.get("agent", "umberto"),
                "deleted_at": None,
                "created_at": now_utc(),
            })
        if docs:
            await db.clients.insert_many(docs)
        assigned = sum(1 for d in docs if d["giro_id"])
        logger.info("Imported %d clients (%d assigned, %d da verificare, %d deleted, %d extra)",
                    len(docs), assigned, len(docs) - assigned, deleted, len(seed_data.EXTRA_CLIENTS))

    await seed_recurrences()


def _norm(s) -> str:
    if s is None:
        return ""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    s = s.lower().strip()
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


async def seed_recurrences():
    # Seed recurrence definitions (idempotent by company key).
    for d in recurrence_seed.RECURRENCE_DEFS:
        existing = await db.recurrence_defs.find_one({"company": d["company"]})
        if not existing:
            await db.recurrence_defs.insert_one({
                "id": str(uuid.uuid4()),
                "company": d["company"],
                "label": d["label"],
                "periods": d["periods"],
                "groups": d["groups"],
                "order": d["order"],
                "created_at": now_utc(),
            })

    # Import members once (guard marker).
    marker = await db.app_meta.find_one({"key": "recurrences_seeded_v1"})
    if marker:
        return

    # Build lookup of existing active clients by normalized name(+city).
    existing = await db.clients.find({"deleted_at": None}, {"_id": 0, "id": 1, "ragione_sociale": 1, "citta": 1, "agent": 1}).to_list(20000)
    by_namecity = {}
    for c in existing:
        key = _norm(c.get("ragione_sociale")) + "|" + _norm(c.get("citta"))
        by_namecity.setdefault(key, c)

    total, matched, created = 0, 0, 0
    for company, members in recurrence_seed.RECURRENCE_MEMBERS.items():
        pos_by_group = {}
        for m in members:
            total += 1
            key = _norm(m["ragione_sociale"]) + "|" + _norm(m["citta"])
            match = by_namecity.get(key)
            if match:
                client_id = match["id"]
                agent = match.get("agent") or m["agent"]
                matched += 1
            else:
                # Recurrence-only new client: no territorial giro, flagged for review
                # unless explicitly declared new in the prompt.
                client_id = str(uuid.uuid4())
                agent = m["agent"]
                await db.clients.insert_one({
                    "id": client_id,
                    "ragione_sociale": m["ragione_sociale"],
                    "codice_azienda": m.get("codice_azienda", ""),
                    "provincia": m.get("provincia", ""),
                    "giro_id": None,
                    "position": 999,
                    "citta": m.get("citta", ""),
                    "zona": m.get("citta", ""),
                    "indirizzo": m.get("indirizzo", ""),
                    "cap": m.get("cap", ""),
                    "telefono": m.get("telefono", ""),
                    "email": m.get("email", ""),
                    "agent": agent,
                    "permanent_note": "",
                    "last_visit_at": None,
                    "snoozed_until": None,
                    "extra": {
                        "recurrence_only": True,
                        "needs_review": not m.get("explicit_new", False),
                        "partita_iva": m.get("partita_iva", ""),
                        "codice_fiscale": m.get("codice_fiscale", ""),
                    },
                    "created_by": agent,
                    "deleted_at": None,
                    "created_at": now_utc(),
                })
                created += 1
            g = m["group"]
            pos = pos_by_group.get(g, 0)
            pos_by_group[g] = pos + 1
            # Avoid duplicate membership for same (company, client).
            dup = await db.recurrence_members.find_one({"company": company, "client_id": client_id})
            if not dup:
                await db.recurrence_members.insert_one({
                    "id": str(uuid.uuid4()),
                    "company": company,
                    "client_id": client_id,
                    "group": g,
                    "position": pos,
                    "agent": agent,
                    "created_at": now_utc(),
                })

    await db.app_meta.insert_one({"key": "recurrences_seeded_v1", "at": now_utc()})
    logger.info("Recurrences seeded: %d members (%d matched existing, %d created new)", total, matched, created)


@app.on_event("startup")
async def on_startup():
    await seed()


@app.on_event("shutdown")
async def on_shutdown():
    client.close()


# ---------------------------------------------------------------------------
# Auth routes
# ---------------------------------------------------------------------------
@api.get("/")
async def root():
    return {"message": "AgendaVisite API"}


@api.post("/auth/login")
async def login(body: LoginRequest):
    username = body.username.strip().lower()
    user = await db.users.find_one({"username": username})
    dummy = "$2b$12$" + "x" * 53
    if not user:
        verify_pw(body.password, dummy)
        raise HTTPException(status_code=401, detail="Username o password errati")
    if not user.get("is_active", False) or not verify_pw(body.password, user["hashed_password"]):
        raise HTTPException(status_code=401, detail="Username o password errati")
    token = create_token(user["username"])
    return {
        "access_token": token,
        "token_type": "bearer",
        "user": {"username": user["username"], "display_name": user["display_name"], "role": user.get("role", "agent")},
    }


@api.get("/auth/me")
async def me(user=Depends(get_current_user)):
    return {"username": user["username"], "display_name": user["display_name"], "role": user.get("role", "agent")}


async def require_admin(user=Depends(get_current_user)):
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Solo l'amministratore può eseguire questa operazione")
    return user


# ---------------------------------------------------------------------------
# Companies
# ---------------------------------------------------------------------------
@api.get("/companies")
async def list_companies(include_inactive: bool = False, user=Depends(get_current_user)):
    q = {} if include_inactive else {"active": True}
    docs = await db.companies.find(q, {"_id": 0}).sort("order", 1).to_list(1000)
    return docs


@api.post("/companies")
async def create_company(body: Company, user=Depends(require_admin)):
    name = body.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Nome azienda obbligatorio")
    existing = await db.companies.find_one({"name": name})
    if existing:
        raise HTTPException(status_code=400, detail="Azienda già esistente")
    last = await db.companies.find_one({}, {"_id": 0}, sort=[("order", -1)])
    order = (last["order"] + 1) if last else 0
    doc = {"id": str(uuid.uuid4()), "name": name, "active": True, "order": order, "created_at": now_utc()}
    await db.companies.insert_one(doc)
    doc.pop("_id", None)
    doc.pop("created_at", None)
    return doc


@api.put("/companies/{company_id}")
async def update_company(company_id: str, body: CompanyUpdate, user=Depends(require_admin)):
    update = {k: v for k, v in body.dict().items() if v is not None}
    if not update:
        raise HTTPException(status_code=400, detail="Nessun dato")
    res = await db.companies.update_one({"id": company_id}, {"$set": update})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Azienda non trovata")
    doc = await db.companies.find_one({"id": company_id}, {"_id": 0})
    return doc


# ---------------------------------------------------------------------------
# Giri
# ---------------------------------------------------------------------------
@api.get("/giri")
async def list_giri(include_inactive: bool = False, user=Depends(get_current_user)):
    q = {} if include_inactive else {"active": True}
    docs = await db.giri.find(q, {"_id": 0}).sort("order", 1).to_list(1000)
    return docs


@api.post("/giri")
async def create_giro(body: GiroCreate, user=Depends(require_admin)):
    name = body.name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Nome giro obbligatorio")
    last = await db.giri.find_one({}, {"_id": 0}, sort=[("order", -1)])
    order = (last["order"] + 1) if last else 0
    doc = {"id": str(uuid.uuid4()), "name": name, "localities": body.localities,
           "order": order, "active": True, "created_at": now_utc()}
    await db.giri.insert_one(doc)
    doc.pop("_id", None)
    doc.pop("created_at", None)
    return doc


@api.put("/giri/{giro_id}")
async def update_giro(giro_id: str, body: GiroUpdate, user=Depends(require_admin)):
    update = {k: v for k, v in body.dict().items() if v is not None}
    if not update:
        raise HTTPException(status_code=400, detail="Nessun dato")
    res = await db.giri.update_one({"id": giro_id}, {"$set": update})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Giro non trovato")
    doc = await db.giri.find_one({"id": giro_id}, {"_id": 0})
    return doc


# ---------------------------------------------------------------------------
# Clients
# ---------------------------------------------------------------------------
async def _handled_today_ids(client_ids: List[str]) -> set:
    if not client_ids:
        return set()
    start = start_of_today_utc()
    cursor = db.events.find(
        {"client_id": {"$in": client_ids}, "created_at": {"$gte": start}, "deleted_at": None,
         "type": {"$nin": ["suspension"]}},
        {"_id": 0, "client_id": 1},
    )
    ids = set()
    async for e in cursor:
        ids.add(e["client_id"])
    return ids


async def _active_suspensions_for(client_ids: List[str]) -> dict:
    """Return {client_id: [company_name,...]} of currently active suspensions.
    A suspension (manual event or overdue deferred order) for a company is active
    unless a later 'collection' (incassato) for that same company settled it.
    """
    if not client_ids:
        return {}
    now = now_utc()
    events = await db.events.find(
        {"client_id": {"$in": client_ids}, "deleted_at": None,
         "type": {"$in": ["order", "suspension", "collection"]}},
        {"_id": 0},
    ).to_list(200000)

    from collections import defaultdict

    def _aware(dt):
        if dt is None:
            return None
        return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)

    data = defaultdict(lambda: defaultdict(lambda: {"paid": None, "manual": [], "auto": []}))
    for e in events:
        comp = e.get("company_name") or e.get("company_id")
        if not comp:
            continue
        cid = e["client_id"]
        created = _aware(e.get("created_at"))
        slot = data[cid][comp]
        if e["type"] == "collection":
            if slot["paid"] is None or (created and created > slot["paid"]):
                slot["paid"] = created
        elif e["type"] == "suspension":
            slot["manual"].append(created)
        elif e["type"] == "order":
            due = _aware(e.get("due_at"))
            if due is not None and due <= now:
                slot["auto"].append(created)

    result = {}
    for cid, comps in data.items():
        active = []
        for comp, slot in comps.items():
            paid = slot["paid"]
            man = any(paid is None or (c and c > paid) for c in slot["manual"])
            auto = any(paid is None or (c and c > paid) for c in slot["auto"])
            if man or auto:
                active.append(comp)
        if active:
            result[cid] = sorted(active, key=lambda x: x.lower())
    return result


def _compute_status(doc: dict, handled_today: set) -> str:
    now = now_utc()
    if doc["id"] in handled_today:
        return "gestito"
    snoozed = doc.get("snoozed_until")
    if snoozed is not None:
        if snoozed.tzinfo is None:
            snoozed = snoozed.replace(tzinfo=timezone.utc)
        if snoozed > now:
            return "gestito"
    lv = doc.get("last_visit_at")
    if lv is None:
        return "da_visitare"
    if lv.tzinfo is None:
        lv = lv.replace(tzinfo=timezone.utc)
    if lv <= now - timedelta(days=VISIT_THRESHOLD_DAYS):
        return "da_visitare"
    return "gestito"


@api.get("/clients")
async def list_clients(giro_id: str = Query(...), user=Depends(get_current_user)):
    docs = await db.clients.find(
        {"giro_id": giro_id, "agent": user["username"], "deleted_at": None}, {"_id": 0}
    ).to_list(5000)
    docs.sort(key=lambda d: (d.get("position", 999), d.get("ragione_sociale", "").lower()))
    ids = [d["id"] for d in docs]
    handled = await _handled_today_ids(ids)
    suspensions = await _active_suspensions_for(ids)
    result = []
    for d in docs:
        pub = client_public(d)
        pub["status"] = _compute_status(d, handled)
        pub["handled_today"] = d["id"] in handled
        pub["suspensions"] = suspensions.get(d["id"], [])
        result.append(pub)
    return result


@api.get("/clients/da-verificare")
async def da_verificare(user=Depends(get_current_user)):
    q = {"deleted_at": None, "$or": [{"giro_id": None}, {"extra.needs_review": True}]}
    if user.get("role") != "admin":
        q["agent"] = user["username"]
    docs = await db.clients.find(q, {"_id": 0}).to_list(5000)
    docs.sort(key=lambda d: d.get("ragione_sociale", "").lower())
    return [client_public(d) for d in docs]


VALID_AGENTS = {"umberto", "andrea"}


@api.get("/clients/all")
async def list_all_clients(search: str = Query("", alias="search"), user=Depends(get_current_user)):
    q = {"deleted_at": None}
    if user.get("role") != "admin":
        q["agent"] = user["username"]
    s = (search or "").strip()
    if s:
        q["ragione_sociale"] = {"$regex": re.escape(s), "$options": "i"}
    docs = await db.clients.find(q, {"_id": 0}).to_list(10000)
    docs.sort(key=lambda d: d.get("ragione_sociale", "").lower())
    return [client_public(d) for d in docs]


@api.post("/clients")
async def create_client(body: ClientCreate, user=Depends(get_current_user)):
    if not body.ragione_sociale.strip():
        raise HTTPException(status_code=400, detail="Ragione sociale obbligatoria")
    agent = user["username"]
    if user.get("role") == "admin" and body.agent in VALID_AGENTS:
        agent = body.agent
    doc = {
        "id": str(uuid.uuid4()),
        "ragione_sociale": body.ragione_sociale.strip(),
        "codice_azienda": "",
        "provincia": body.provincia.strip(),
        "giro_id": body.giro_id,
        "position": body.position if body.position is not None else 999,
        "citta": body.citta.strip(),
        "zona": body.zona.strip(),
        "indirizzo": body.indirizzo.strip(),
        "cap": body.cap.strip(),
        "telefono": body.telefono.strip(),
        "email": body.email.strip(),
        "agent": agent,
        "permanent_note": "",
        "last_visit_at": None,
        "snoozed_until": None,
        "extra": {},
        "created_by": user["username"],
        "deleted_at": None,
        "created_at": now_utc(),
    }
    await db.clients.insert_one(doc)
    return client_public(doc)


@api.get("/clients/{client_id}")
async def get_client(client_id: str, user=Depends(get_current_user)):
    doc = await db.clients.find_one({"id": client_id, "deleted_at": None}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    pub = client_public(doc)
    susp = await _active_suspensions_for([client_id])
    pub["suspensions"] = susp.get(client_id, [])
    return pub


@api.put("/clients/{client_id}")
async def update_client(client_id: str, body: ClientUpdate, user=Depends(get_current_user)):
    update = {k: v for k, v in body.dict(exclude_unset=True).items()}
    if not update:
        raise HTTPException(status_code=400, detail="Nessun dato")
    target = await db.clients.find_one({"id": client_id, "deleted_at": None}, {"_id": 0})
    if not target:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    if user.get("role") != "admin" and target.get("agent") != user["username"]:
        raise HTTPException(status_code=403, detail="Non puoi modificare questo cliente")
    # Only admin may (re)assign the agent; validate value.
    if "agent" in update:
        if user.get("role") != "admin" or update["agent"] not in VALID_AGENTS:
            update.pop("agent", None)
    set_doc = dict(update)
    if update.get("giro_id"):
        set_doc["extra.needs_review"] = False
    await db.clients.update_one({"id": client_id, "deleted_at": None}, {"$set": set_doc})
    doc = await db.clients.find_one({"id": client_id}, {"_id": 0})
    return client_public(doc)


@api.delete("/clients/{client_id}")
async def delete_client(client_id: str, user=Depends(require_admin)):
    res = await db.clients.update_one(
        {"id": client_id, "deleted_at": None}, {"$set": {"deleted_at": now_utc()}}
    )
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    return {"ok": True}


@api.get("/clients/{client_id}/history")
async def client_history(client_id: str, user=Depends(get_current_user)):
    docs = await db.events.find(
        {"client_id": client_id, "deleted_at": None}, {"_id": 0}
    ).sort("created_at", -1).to_list(5000)
    return [event_public(d) for d in docs]


# ---------------------------------------------------------------------------
# Events (quick actions)
# ---------------------------------------------------------------------------
VALID_EVENT_TYPES = {"visit", "order", "reschedule", "collection", "note", "suspension"}


@api.get("/payment-modes")
async def payment_modes(user=Depends(get_current_user)):
    return [{"key": k, "label": PAYMENT_MODES[k]["label"], "days": PAYMENT_MODES[k]["days"]} for k in PAYMENT_MODE_ORDER]


@api.post("/events")
async def create_event(body: EventCreate, user=Depends(get_current_user)):
    if body.type not in VALID_EVENT_TYPES:
        raise HTTPException(status_code=400, detail="Tipo evento non valido")
    cli = await db.clients.find_one({"id": body.client_id, "deleted_at": None}, {"_id": 0})
    if not cli:
        raise HTTPException(status_code=404, detail="Cliente non trovato")

    now = now_utc()
    doc = {
        "id": str(uuid.uuid4()),
        "client_id": body.client_id,
        "type": body.type,
        "company_id": None,
        "company_name": None,
        "note_text": None,
        "reschedule_until": None,
        "payment_mode": None,
        "due_at": None,
        "source": None,
        "agent": user["username"],
        "created_at": now,
        "deleted_at": None,
    }

    if body.type in ("order", "collection", "suspension"):
        if not body.company_id:
            raise HTTPException(status_code=400, detail="Seleziona un'azienda")
        comp = await db.companies.find_one({"id": body.company_id}, {"_id": 0})
        if not comp:
            raise HTTPException(status_code=404, detail="Azienda non trovata")
        doc["company_id"] = comp["id"]
        doc["company_name"] = comp["name"]

    if body.type == "order":
        # Payment mode chosen for THIS order only (does not touch client anagrafica).
        if body.payment_mode:
            if body.payment_mode not in PAYMENT_MODES:
                raise HTTPException(status_code=400, detail="Modalità di pagamento non valida")
            doc["payment_mode"] = body.payment_mode
            days = PAYMENT_MODES[body.payment_mode]["days"]
            if days is not None:
                doc["due_at"] = now + timedelta(days=days)

    if body.type == "suspension":
        doc["source"] = "manual"

    if body.type == "note":
        if not (body.note_text and body.note_text.strip()):
            raise HTTPException(status_code=400, detail="Nota vuota")
        doc["note_text"] = body.note_text.strip()

    if body.type == "reschedule":
        target = None
        if body.reschedule_date:
            try:
                target = datetime.fromisoformat(body.reschedule_date.replace("Z", "+00:00"))
                if target.tzinfo is None:
                    target = target.replace(tzinfo=timezone.utc)
            except ValueError:
                raise HTTPException(status_code=400, detail="Data non valida")
        elif body.reschedule_days:
            target = now + timedelta(days=body.reschedule_days)
        doc["reschedule_until"] = target
        if target is not None:
            await db.clients.update_one({"id": body.client_id}, {"$set": {"snoozed_until": target}})

    await db.events.insert_one(doc)

    if body.type == "visit":
        await db.clients.update_one(
            {"id": body.client_id},
            {"$set": {"last_visit_at": now, "snoozed_until": None}},
        )

    return event_public(doc)


# ---------------------------------------------------------------------------
# Activities feed (Ultimi aggiornamenti / Storico) — read only
# ---------------------------------------------------------------------------
ACTIVITY_TYPES = ["order", "collection", "suspension", "reschedule"]
ACTIVITY_TYPE_LABELS = {
    "order": "Ordine effettuato",
    "collection": "Incassato",
    "suspension": "+Sospeso",
    "reschedule": "Visita rimandata",
}


@api.get("/activities")
async def list_activities(
    scope: str = Query("all"),  # all | umberto | andrea
    type: str = Query("all"),   # all | order | collection | suspension | reschedule
    limit: int = Query(200),
    user=Depends(get_current_user),
):
    is_admin = user.get("role") == "admin"
    q = {"deleted_at": None, "type": {"$in": ACTIVITY_TYPES}}
    if type != "all" and type in ACTIVITY_TYPES:
        q["type"] = type
    # Permissions: agents only ever see their own activities regardless of scope.
    if not is_admin:
        q["agent"] = user["username"]
    elif scope in VALID_AGENTS:
        q["agent"] = scope
    lim = max(1, min(limit, 500))
    events = await db.events.find(q, {"_id": 0}).sort("created_at", -1).limit(lim).to_list(lim)

    client_ids = list({e["client_id"] for e in events})
    clients = await db.clients.find({"id": {"$in": client_ids}}, {"_id": 0}).to_list(20000)
    cmap = {c["id"]: c for c in clients}
    giri = await db.giri.find({}, {"_id": 0, "id": 1, "name": 1}).to_list(1000)
    gmap = {g["id"]: g["name"] for g in giri}

    out = []
    for e in events:
        cli = cmap.get(e["client_id"], {})
        giro_name = gmap.get(cli.get("giro_id")) if cli else None
        if e["type"] in ("order", "collection", "suspension"):
            context = e.get("company_name") or ""
        else:
            context = giro_name or (cli.get("citta") if cli else "") or ""
        out.append({
            "id": e["id"],
            "type": e["type"],
            "type_label": ACTIVITY_TYPE_LABELS.get(e["type"], e["type"]),
            "created_at": iso(e.get("created_at")),
            "agent": e.get("agent"),
            "client_ragione_sociale": cli.get("ragione_sociale", "") if cli else "",
            "context": context,
            "giro_name": giro_name,
            "citta": cli.get("citta", "") if cli else "",
            "company_name": e.get("company_name"),
        })
    return out


# ---------------------------------------------------------------------------
# Recurrences (Ricorrenze) — separate from territorial giri
# ---------------------------------------------------------------------------
class RecurrenceOrderBody(BaseModel):
    client_id: str
    period: str


class RecurrenceMemberCreate(BaseModel):
    group: str
    period: Optional[str] = None
    client_id: Optional[str] = None
    ragione_sociale: Optional[str] = None
    citta: Optional[str] = ""
    provincia: Optional[str] = ""
    indirizzo: Optional[str] = ""
    cap: Optional[str] = ""
    telefono: Optional[str] = ""
    email: Optional[str] = ""
    agent: Optional[str] = None


def _rome_year_start() -> datetime:
    now_rome = now_utc().astimezone(ROME)
    return datetime(now_rome.year, 1, 1, tzinfo=ROME).astimezone(timezone.utc)


async def _recurrence_done_ids(company: str, period: str, client_ids: List[str]) -> dict:
    """Return {client_id: order_iso_date} for orders placed this calendar year."""
    if not client_ids:
        return {}
    start = _rome_year_start()
    cursor = db.events.find(
        {"type": "recurrence_order", "recurrence_company": company,
         "recurrence_period": period, "client_id": {"$in": client_ids},
         "created_at": {"$gte": start}, "deleted_at": None},
        {"_id": 0, "client_id": 1, "created_at": 1},
    )
    out = {}
    async for e in cursor:
        cid = e["client_id"]
        if cid not in out:
            out[cid] = iso(e["created_at"])
    return out


async def _get_recur_def(company: str) -> dict:
    d = await db.recurrence_defs.find_one({"company": company}, {"_id": 0})
    if not d:
        raise HTTPException(status_code=404, detail="Ricorrenza non trovata")
    return d


@api.get("/recurrences")
async def list_recurrences(user=Depends(get_current_user)):
    defs = await db.recurrence_defs.find({}, {"_id": 0}).sort("order", 1).to_list(100)
    return defs


@api.get("/recurrences/{company}/members")
async def recurrence_members(company: str, period: str = Query(...), user=Depends(get_current_user)):
    rdef = await _get_recur_def(company)
    if period not in [p["key"] for p in rdef["periods"]]:
        raise HTTPException(status_code=400, detail="Periodo non valido")

    members = await db.recurrence_members.find({"company": company}, {"_id": 0}).to_list(20000)
    client_ids = [m["client_id"] for m in members]
    clients = await db.clients.find(
        {"id": {"$in": client_ids}, "deleted_at": None}, {"_id": 0}
    ).to_list(20000)
    cmap = {c["id"]: c for c in clients}

    is_admin = user.get("role") == "admin"
    visible = []
    for m in members:
        cli = cmap.get(m["client_id"])
        if not cli:
            continue  # client deleted
        if not is_admin and cli.get("agent") != user["username"]:
            continue
        visible.append((m, cli))

    done = await _recurrence_done_ids(company, period, [c["id"] for _, c in visible])

    # Group by recurrence group, preserving def group order then position.
    group_order = {g: i for i, g in enumerate(rdef["groups"])}
    visible.sort(key=lambda mc: (group_order.get(mc[0]["group"], 999),
                                 mc[0].get("position", 999),
                                 mc[1].get("ragione_sociale", "").lower()))
    groups_out = []
    current = None
    for m, cli in visible:
        if current is None or current["group"] != m["group"]:
            current = {"group": m["group"], "clients": []}
            groups_out.append(current)
        pub = client_public(cli)
        pub["recurrence_status"] = "ordine_effettuato" if cli["id"] in done else "da_gestire"
        pub["order_date"] = done.get(cli["id"])
        pub["member_id"] = m["id"]
        current["clients"].append(pub)

    return {"company": company, "period": period, "groups": groups_out}


@api.post("/recurrences/{company}/order")
async def recurrence_order(company: str, body: RecurrenceOrderBody, user=Depends(get_current_user)):
    rdef = await _get_recur_def(company)
    if body.period not in [p["key"] for p in rdef["periods"]]:
        raise HTTPException(status_code=400, detail="Periodo non valido")
    cli = await db.clients.find_one({"id": body.client_id, "deleted_at": None}, {"_id": 0})
    if not cli:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    if user.get("role") != "admin" and cli.get("agent") != user["username"]:
        raise HTTPException(status_code=403, detail="Non puoi gestire questo cliente")
    now = now_utc()
    doc = {
        "id": str(uuid.uuid4()),
        "client_id": body.client_id,
        "type": "recurrence_order",
        "company_id": None,
        "company_name": None,
        "note_text": None,
        "reschedule_until": None,
        "recurrence_company": company,
        "recurrence_period": body.period,
        "agent": user["username"],
        "created_at": now,
        "deleted_at": None,
    }
    await db.events.insert_one(doc)
    return {"ok": True, "order_date": iso(now)}


@api.post("/recurrences/{company}/order/undo")
async def recurrence_order_undo(company: str, body: RecurrenceOrderBody, user=Depends(get_current_user)):
    cli = await db.clients.find_one({"id": body.client_id, "deleted_at": None}, {"_id": 0})
    if not cli:
        raise HTTPException(status_code=404, detail="Cliente non trovato")
    if user.get("role") != "admin" and cli.get("agent") != user["username"]:
        raise HTTPException(status_code=403, detail="Non puoi gestire questo cliente")
    start = _rome_year_start()
    await db.events.update_many(
        {"type": "recurrence_order", "recurrence_company": company,
         "recurrence_period": body.period, "client_id": body.client_id,
         "created_at": {"$gte": start}, "deleted_at": None},
        {"$set": {"deleted_at": now_utc()}},
    )
    return {"ok": True}


@api.post("/recurrences/{company}/members")
async def add_recurrence_member(company: str, body: RecurrenceMemberCreate, user=Depends(get_current_user)):
    rdef = await _get_recur_def(company)
    if body.group not in rdef["groups"]:
        raise HTTPException(status_code=400, detail="Gruppo non valido")

    agent = user["username"]
    if user.get("role") == "admin" and body.agent in VALID_AGENTS:
        agent = body.agent

    if body.client_id:
        cli = await db.clients.find_one({"id": body.client_id, "deleted_at": None}, {"_id": 0})
        if not cli:
            raise HTTPException(status_code=404, detail="Cliente non trovato")
        if user.get("role") != "admin" and cli.get("agent") != user["username"]:
            raise HTTPException(status_code=403, detail="Non puoi aggiungere questo cliente")
        client_id = cli["id"]
        agent = cli.get("agent") or agent
    else:
        if not (body.ragione_sociale and body.ragione_sociale.strip()):
            raise HTTPException(status_code=400, detail="Ragione sociale obbligatoria")
        client_id = str(uuid.uuid4())
        await db.clients.insert_one({
            "id": client_id,
            "ragione_sociale": body.ragione_sociale.strip(),
            "codice_azienda": "",
            "provincia": (body.provincia or "").strip(),
            "giro_id": None,
            "position": 999,
            "citta": (body.citta or "").strip(),
            "zona": (body.citta or "").strip(),
            "indirizzo": (body.indirizzo or "").strip(),
            "cap": (body.cap or "").strip(),
            "telefono": (body.telefono or "").strip(),
            "email": (body.email or "").strip(),
            "agent": agent,
            "permanent_note": "",
            "last_visit_at": None,
            "snoozed_until": None,
            "extra": {"recurrence_only": True, "needs_review": False},
            "created_by": user["username"],
            "deleted_at": None,
            "created_at": now_utc(),
        })

    dup = await db.recurrence_members.find_one({"company": company, "client_id": client_id})
    if dup:
        raise HTTPException(status_code=400, detail="Cliente già presente in questa ricorrenza")

    last = await db.recurrence_members.find(
        {"company": company, "group": body.group}, {"_id": 0, "position": 1}
    ).sort("position", -1).to_list(1)
    pos = (last[0]["position"] + 1) if last else 0
    await db.recurrence_members.insert_one({
        "id": str(uuid.uuid4()),
        "company": company,
        "client_id": client_id,
        "group": body.group,
        "position": pos,
        "agent": agent,
        "created_at": now_utc(),
    })
    return {"ok": True, "client_id": client_id}


@api.delete("/recurrences/{company}/members/{member_id}")
async def remove_recurrence_member(company: str, member_id: str, user=Depends(get_current_user)):
    member = await db.recurrence_members.find_one({"id": member_id, "company": company}, {"_id": 0})
    if not member:
        raise HTTPException(status_code=404, detail="Voce non trovata")
    if user.get("role") != "admin":
        cli = await db.clients.find_one({"id": member["client_id"]}, {"_id": 0})
        if not cli or cli.get("agent") != user["username"]:
            raise HTTPException(status_code=403, detail="Non puoi rimuovere questo cliente")
    await db.recurrence_members.delete_one({"id": member_id, "company": company})
    return {"ok": True}


# ---------------------------------------------------------------------------
# Monthly statistics + Excel export (per giro)
# ---------------------------------------------------------------------------
MONTHS_IT = ["Gen", "Feb", "Mar", "Apr", "Mag", "Giu", "Lug", "Ago", "Set", "Ott", "Nov", "Dic"]


async def _monthly_grid(giro_id: str, year: int, user: dict):
    giro = await db.giri.find_one({"id": giro_id}, {"_id": 0})
    if not giro:
        raise HTTPException(status_code=404, detail="Giro non trovato")

    q = {"giro_id": giro_id, "deleted_at": None}
    if user.get("role") != "admin":
        q["agent"] = user["username"]
    clients = await db.clients.find(q, {"_id": 0}).to_list(5000)
    clients.sort(key=lambda d: (d.get("position", 999), d.get("ragione_sociale", "").lower()))
    client_ids = [c["id"] for c in clients]

    year_start = datetime(year, 1, 1, tzinfo=timezone.utc)
    year_end = datetime(year + 1, 1, 1, tzinfo=timezone.utc)
    events = await db.events.find(
        {"client_id": {"$in": client_ids}, "deleted_at": None,
         "created_at": {"$gte": year_start, "$lt": year_end}},
        {"_id": 0},
    ).to_list(100000)

    grid = {cid: {m: {"visit": False, "orders": [], "collection": False} for m in range(12)} for cid in client_ids}
    for e in events:
        created = e["created_at"]
        if created.tzinfo is None:
            created = created.replace(tzinfo=timezone.utc)
        m = created.astimezone(ROME).month - 1
        cell = grid[e["client_id"]][m]
        if e["type"] == "visit":
            cell["visit"] = True
        elif e["type"] == "order":
            name = e.get("company_name") or "Ordine"
            if name not in cell["orders"]:
                cell["orders"].append(name)
        elif e["type"] == "collection":
            cell["collection"] = True
    return giro, clients, grid


@api.get("/stats/monthly")
async def stats_monthly(giro_id: str = Query(...), year: int = Query(...), user=Depends(get_current_user)):
    giro, clients, grid = await _monthly_grid(giro_id, year, user)
    rows = []
    for cli in clients:
        months = []
        for m in range(12):
            c = grid[cli["id"]][m]
            months.append({"visit": c["visit"], "orders": c["orders"], "collection": c["collection"]})
        rows.append({
            "id": cli["id"],
            "ragione_sociale": cli.get("ragione_sociale", ""),
            "citta": cli.get("citta", ""),
            "months": months,
        })
    return {"giro": {"id": giro["id"], "name": giro["name"]}, "year": year, "months": MONTHS_IT, "clients": rows}


@api.get("/export/monthly")
async def export_monthly(giro_id: str = Query(...), year: int = Query(...), user=Depends(get_current_user)):
    giro, clients, grid = await _monthly_grid(giro_id, year, user)

    def mark_text(cell: dict) -> str:
        out = []
        if cell["visit"]:
            out.append("\u2713 visita")
        if cell["orders"]:
            out.append("\u2713 ordine: " + ", ".join(cell["orders"]))
        if cell["collection"]:
            out.append("\u2713 incasso")
        return "\n".join(out)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Riepilogo"

    header_fill = PatternFill("solid", fgColor="047857")
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14)

    ws.cell(row=1, column=1, value=f"Giro: {giro['name']} \u2014 Anno {year}").font = title_font
    ws.append([])

    headers = ["Cliente", "Citt\u00e0"] + MONTHS_IT
    ws.append(headers)
    hrow = ws.max_row
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=hrow, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    for cli in clients:
        row = [cli.get("ragione_sociale", ""), cli.get("citta", "")]
        for m in range(12):
            row.append(mark_text(grid[cli["id"]][m]))
        ws.append(row)
        r = ws.max_row
        for col in range(3, 15):
            ws.cell(row=r, column=col).alignment = Alignment(horizontal="center", wrap_text=True, vertical="top")

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 20
    for i in range(12):
        ws.column_dimensions[openpyxl.utils.get_column_letter(3 + i)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    safe = "".join(ch for ch in giro["name"] if ch.isalnum() or ch in " -_").strip().replace(" ", "_")
    filename = f"Riepilogo_{safe}_{year}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


app.include_router(api)
app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)
