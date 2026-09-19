import re
import unicodedata
import openpyxl
from pymongo import MongoClient

client = MongoClient("mongodb://localhost:27017")
db = client["test_database"]

def norm(s):
    if s is None:
        return ""
    s = str(s)
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    s = s.lower().strip()
    s = re.sub(r"[^a-z0-9]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s

# build lookup of existing clients
existing = list(db.clients.find({"deleted_at": None}, {"_id": 0, "id": 1, "ragione_sociale": 1, "citta": 1, "agent": 1, "giro_id": 1}))
by_name = {}
by_namecity = {}
for c in existing:
    n = norm(c.get("ragione_sociale"))
    nc = n + "|" + norm(c.get("citta"))
    by_name.setdefault(n, []).append(c)
    by_namecity.setdefault(nc, []).append(c)

print("Existing active clients:", len(existing))

files = {
    "mazzetti": "/app/data/ric/mazzetti.xlsx",
    "bonfissuto": "/app/data/ric/bonfissuto.xlsx",
}

for fname, path in files.items():
    print("="*60)
    print("FILE:", fname)
    wb = openpyxl.load_workbook(path, data_only=True)
    total = 0
    matched = 0
    matched_name_only = 0
    unmatched = []
    seen = set()
    for ws in wb.worksheets:
        for r in range(2, ws.max_row + 1):
            rag = ws.cell(row=r, column=2).value
            citta = ws.cell(row=r, column=5).value
            if not rag or not str(rag).strip():
                continue
            key = norm(rag) + "|" + norm(citta)
            if key in seen:
                continue
            seen.add(key)
            total += 1
            nc = norm(rag) + "|" + norm(citta)
            n = norm(rag)
            if nc in by_namecity:
                matched += 1
            elif n in by_name:
                matched_name_only += 1
            else:
                unmatched.append(f"{rag} | {citta} [{ws.title}]")
    print(f"unique rows: {total} | matched name+city: {matched} | matched name-only(diff city): {matched_name_only} | UNMATCHED: {len(unmatched)}")
    print("--- UNMATCHED (would be NEW recurrence-only clients) ---")
    for u in unmatched:
        print("  ", u)
