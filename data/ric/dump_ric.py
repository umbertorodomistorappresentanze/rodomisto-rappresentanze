import openpyxl

for f in ["mazzetti", "bonfissuto"]:
    print("=" * 70)
    print("FILE:", f)
    wb = openpyxl.load_workbook(f"/app/data/ric/{f}.xlsx", data_only=True)
    for ws in wb.worksheets:
        print(f"--- SHEET: {ws.title} | dims={ws.dimensions} | max_row={ws.max_row} max_col={ws.max_column}")
        # print first 6 rows to detect header
        for r in range(1, min(ws.max_row, 6) + 1):
            vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            print(r, vals)
