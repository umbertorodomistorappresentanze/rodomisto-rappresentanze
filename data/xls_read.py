import openpyxl
wb = openpyxl.load_workbook('/app/data/clienti.xlsx', data_only=True)
for ws in wb.worksheets:
    print("=== SHEET:", ws.title, "dims:", ws.dimensions, "max_row:", ws.max_row, "max_col:", ws.max_column)
    # header row
    rows = list(ws.iter_rows(values_only=True))
    print("HEADER:", rows[0])
    print("N data rows (excl header):", len(rows)-1)
    print("--- first 6 data rows ---")
    for r in rows[1:7]:
        print(r)
