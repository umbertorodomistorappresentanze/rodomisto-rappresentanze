import sys, unicodedata, re
sys.path.insert(0, "/app/backend")
import openpyxl
from pymongo import MongoClient
import seed_data as sd

db = MongoClient("mongodb://localhost:27017")["test_database"]

def pk(s):  # piva/cf normalize
    if s is None: return ""
    return re.sub(r"[^0-9a-zA-Z]", "", str(s)).lower()

existing = list(db.clients.find({"deleted_at": None}, {"_id":0,"id":1,"ragione_sociale":1,"citta":1,"indirizzo":1,"agent":1,"giro_id":1,"extra":1}))
by_piva, by_cf, by_nc, by_name = {}, {}, {}, {}
for c in existing:
    ex = c.get("extra") or {}
    if pk(ex.get("partita_iva")): by_piva.setdefault(pk(ex.get("partita_iva")), []).append(c)
    if pk(ex.get("codice_fiscale")): by_cf.setdefault(pk(ex.get("codice_fiscale")), []).append(c)
    n = sd._key(c.get("ragione_sociale"))
    by_nc.setdefault(n+"|"+sd._key(c.get("citta")), []).append(c)
    by_name.setdefault(n, []).append(c)

print("Esistenti attivi:", len(existing))

giri = list(db.giri.find({}, {"_id":0}))
loc_index = sd.build_locality_index(giri)
giro_name = {g["id"]: g["name"] for g in giri}

wb = openpyxl.load_workbook("/app/data/menu/menu.xlsx", data_only=True)
HDR = None
rows = []
for ws in wb.worksheets:
    header = [ (str(ws.cell(row=1,column=c).value).strip() if ws.cell(row=1,column=c).value else "") for c in range(1, ws.max_column+1)]
    idx = {h:i for i,h in enumerate(header)}
    def cv(r, name):
        i = idx.get(name)
        if i is None: return None
        return ws.cell(row=r, column=i+1).value
    for r in range(2, ws.max_row+1):
        rag = ws.cell(row=r, column=idx["Ragione Sociale"]+1).value
        if not rag or not str(rag).strip(): continue
        rows.append({
            "ragione": str(rag).strip(),
            "citta": str(cv(r,"Città") or "").strip(),
            "indirizzo": str(cv(r,"Indirizzo") or "").strip(),
            "zona": str(cv(r,"Descrizione Zona") or "").strip(),
            "piva": pk(cv(r,"Partita IVA")),
            "cf": pk(cv(r,"Codice Fiscale")),
            "agent_raw": str(cv(r,"Nominativo Agente") or "").strip(),
            "sheet": ws.title,
        })

# dedupe within file by name+city
seen=set(); urows=[]
for x in rows:
    k=sd._key(x["ragione"])+"|"+sd._key(x["citta"])
    if k in seen: continue
    seen.add(k); urows.append(x)

matched, new, verify = [], [], []
for x in urows:
    m = None; how=None
    if x["piva"] and x["piva"] in by_piva: m=by_piva[x["piva"]][0]; how="P.IVA"
    elif x["cf"] and x["cf"] in by_cf: m=by_cf[x["cf"]][0]; how="CF"
    elif (sd._key(x["ragione"])+"|"+sd._key(x["citta"])) in by_nc: m=by_nc[sd._key(x["ragione"])+"|"+sd._key(x["citta"])][0]; how="Nome+Comune"
    if m:
        matched.append((x,m,how))
    else:
        # possible name-only (different city) -> uncertain
        if sd._key(x["ragione"]) in by_name:
            verify.append((x, "Nome uguale ma Comune diverso da un cliente esistente ("+by_name[sd._key(x['ragione'])][0].get('citta','')+")"))
            continue
        gid, pos = sd.resolve_assignment(x["ragione"], x["zona"], giri, loc_index)
        if gid:
            new.append((x, gid, pos))
        else:
            verify.append((x, f"Zona '{x['zona']}' non associabile con certezza a un giro"))

print("\n=== RIEPILOGO ===")
print("CLIENTI NEL FILE (unici):", len(urows))
print("GIA' PRESENTI:", len(matched))
print("NUOVI DA CREARE:", len(new))
print("DA VERIFICARE:", len(verify))

print("\n--- GIA' PRESENTI ---")
for x,m,how in matched:
    print(f"  [{how}] {x['ragione']} ({x['citta']}) -> giro attuale: {giro_name.get(m.get('giro_id')) or 'DA VERIFICARE (nessun giro)'}")
print("\n--- NUOVI (con giro assegnato via Descrizione Zona) ---")
for x,gid,pos in new:
    print(f"  {x['ragione']} ({x['citta']}) zona='{x['zona']}' -> {giro_name.get(gid)} pos~{pos}")
print("\n--- DA VERIFICARE ---")
for x,reason in verify:
    print(f"  {x['ragione']} | Comune: {x['citta']} | Zona: '{x['zona']}' | Motivo: {reason}")
